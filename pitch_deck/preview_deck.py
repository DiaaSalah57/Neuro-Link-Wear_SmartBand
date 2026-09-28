#!/usr/bin/env python3
"""
Quick visual preview + overflow linter for a python-pptx generated deck — no PowerPoint needed.

    pip install python-pptx pymupdf pillow
    python preview_deck.py NeuroLink_Wear_Pitch_Deck.pptx            # -> preview/slide_01.png ... + preview.pdf
    python preview_deck.py deck.pptx --out my_folder --dpi 110

It renders the shapes the builder uses (rectangles, rounded rectangles, ovals, right arrows,
connectors, cropped/rounded pictures, solid + linear-gradient fills with alpha, multi-run text)
and prints a warning for every text frame whose wrapped text would not fit its box.
It is an approximation (DejaVu Sans stands in for Calibri, scaled to similar metrics) meant for
layout QA, not a pixel-perfect PowerPoint clone.
"""
from __future__ import annotations

import argparse
import io
import os

import pymupdf
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn

PT = 1 / 12700.0  # EMU -> pt
FONT_DIR = "/usr/share/fonts/truetype/dejavu"
FONTS = {
    (False, False): os.path.join(FONT_DIR, "DejaVuSans.ttf"),
    (True, False): os.path.join(FONT_DIR, "DejaVuSans-Bold.ttf"),
    (False, True): os.path.join(FONT_DIR, "DejaVuSans-Oblique.ttf"),
    (True, True): os.path.join(FONT_DIR, "DejaVuSans-BoldOblique.ttf"),
}
for k, v in list(FONTS.items()):
    if not os.path.exists(v):
        FONTS[k] = FONTS[(k[0], False)] if os.path.exists(FONTS[(k[0], False)]) else FONTS[(False, False)]
WIDTH_FACTOR = 0.88          # DejaVu Sans is ~12% wider than Calibri
LINE_FACTOR = 1.2            # PowerPoint single spacing ≈ 1.2 × font size
_font_cache: dict[str, pymupdf.Font] = {}


def font_for(bold, italic):
    path = FONTS[(bool(bold), bool(italic))]
    if path not in _font_cache:
        _font_cache[path] = pymupdf.Font(fontfile=path)
    return path, _font_cache[path]


def hex_to_rgb01(h):
    return tuple(int(h[i:i + 2], 16) / 255 for i in (0, 2, 4))


def color_of(el):
    """(rgb01, alpha) from an <a:solidFill>/<a:gs> child colour element."""
    c = el.find(qn("a:srgbClr"))
    if c is None:
        return None, 1.0
    a = c.find(qn("a:alpha"))
    return hex_to_rgb01(c.get("val")), (int(a.get("val")) / 100000 if a is not None else 1.0)


def fill_of(spPr):
    solid = spPr.find(qn("a:solidFill"))
    if solid is not None:
        rgb, al = color_of(solid)
        return ("solid", rgb, al)
    grad = spPr.find(qn("a:gradFill"))
    if grad is not None:
        stops = []
        for gs in grad.find(qn("a:gsLst")):
            rgb, al = color_of(gs)
            stops.append((int(gs.get("pos")) / 100000, rgb, al))
        lin = grad.find(qn("a:lin"))
        ang = int(lin.get("ang")) / 60000 if lin is not None else 0
        return ("grad", stops, ang)
    return None


def line_of(spPr):
    ln = spPr.find(qn("a:ln"))
    if ln is None or ln.find(qn("a:noFill")) is not None:
        return None
    solid = ln.find(qn("a:solidFill"))
    if solid is None:
        return None
    rgb, al = color_of(solid)
    w = int(ln.get("w", "9525")) * PT
    return rgb, al, w


def geom_of(spPr):
    g = spPr.find(qn("a:prstGeom"))
    if g is None:
        return "rect", None
    prst = g.get("prst")
    adj = None
    av = g.find(qn("a:avLst"))
    if av is not None:
        for gd in av.findall(qn("a:gd")):
            if gd.get("name") == "adj":
                adj = int(gd.get("fmla").split()[1]) / 100000
    return prst, adj


def draw_gradient(page, rect, stops, ang):
    """Approximate linear gradient with strips (0° = left→right, 90° = top→bottom)."""
    n = 40
    vertical = 45 <= (ang % 180) < 135
    stops = sorted(stops)
    for i in range(n):
        t0, t1 = i / n, (i + 1) / n
        tm = (t0 + t1) / 2
        # interpolate
        lo = max([s for s in stops if s[0] <= tm], default=stops[0], key=lambda s: s[0])
        hi = min([s for s in stops if s[0] >= tm], default=stops[-1], key=lambda s: s[0])
        if hi[0] == lo[0]:
            rgb, al = lo[1], lo[2]
        else:
            f = (tm - lo[0]) / (hi[0] - lo[0])
            rgb = tuple(lo[1][k] + (hi[1][k] - lo[1][k]) * f for k in range(3))
            al = lo[2] + (hi[2] - lo[2]) * f
        if vertical:
            r = pymupdf.Rect(rect.x0, rect.y0 + rect.height * t0, rect.x1, rect.y0 + rect.height * t1 + 0.5)
        else:
            r = pymupdf.Rect(rect.x0 + rect.width * t0, rect.y0, rect.x0 + rect.width * t1 + 0.5, rect.y1)
        if al > 0.005:
            page.draw_rect(r, color=None, fill=rgb, fill_opacity=al)


def draw_shape(page, shape, warnings, slide_no):
    x, y = shape.left * PT, shape.top * PT
    w, h = shape.width * PT, shape.height * PT
    rect = pymupdf.Rect(x, y, x + w, y + h)
    st = shape.shape_type

    if st == MSO_SHAPE_TYPE.PICTURE:
        img = Image.open(io.BytesIO(shape.image.blob)).convert("RGB")
        iw, ih = img.size
        l, r, t, b = shape.crop_left, shape.crop_right, shape.crop_top, shape.crop_bottom
        img = img.crop((int(iw * l), int(ih * t), int(iw * (1 - r)), int(ih * (1 - b))))
        prst, adj = geom_of(shape._element.spPr)
        if prst == "roundRect":
            scale = 4
            img = img.resize((max(1, int(w * scale)), max(1, int(h * scale))))
            mask = Image.new("L", img.size, 0)
            rad = (adj if adj is not None else 0.16667) * min(img.size)
            ImageDraw.Draw(mask).rounded_rectangle((0, 0, img.size[0] - 1, img.size[1] - 1), radius=rad, fill=255)
            rgba = img.convert("RGBA")
            rgba.putalpha(mask)
            buf = io.BytesIO()
            rgba.save(buf, format="PNG")
        else:
            buf = io.BytesIO()
            img.save(buf, format="JPEG", quality=90)
        page.insert_image(rect, stream=buf.getvalue())
        return

    spPr = shape._element.spPr
    if st == MSO_SHAPE_TYPE.LINE or shape._element.tag.endswith("cxnSp"):
        ln = line_of(spPr)
        if ln:
            rgb, al, lw = ln
            page.draw_line(pymupdf.Point(x, y), pymupdf.Point(x + w, y + h), color=rgb, width=max(lw, 0.5), stroke_opacity=al)
        return

    prst, adj = geom_of(spPr)
    fill = fill_of(spPr)
    ln = line_of(spPr)
    stroke = ln[0] if ln else None
    lw = ln[2] if ln else 0
    if fill and fill[0] == "grad":
        draw_gradient(page, rect, fill[1], fill[2])
        fill_rgb, fill_al = None, 1
    else:
        fill_rgb, fill_al = (fill[1], fill[2]) if fill else (None, 1)

    if prst == "ellipse":
        if fill_rgb or stroke:
            page.draw_oval(rect, color=stroke, fill=fill_rgb, width=lw, fill_opacity=fill_al)
    elif prst == "rightArrow":
        pts = [pymupdf.Point(x, y + h * 0.25), pymupdf.Point(x + w * 0.5, y + h * 0.25), pymupdf.Point(x + w * 0.5, y),
               pymupdf.Point(x + w, y + h / 2), pymupdf.Point(x + w * 0.5, y + h), pymupdf.Point(x + w * 0.5, y + h * 0.75),
               pymupdf.Point(x, y + h * 0.75)]
        page.draw_polyline(pts + [pts[0]], color=stroke, fill=fill_rgb, width=lw, fill_opacity=fill_al)
    else:
        if fill_rgb or stroke:
            radius = None
            if prst == "roundRect":
                a = adj if adj is not None else 0.16667
                radius = a * min(w, h) / min(w, h) if min(w, h) else 0  # fraction of the shorter side
                radius = max(0.0, min(0.5, a))
            page.draw_rect(rect, color=stroke, fill=fill_rgb, width=lw, fill_opacity=fill_al, radius=radius)

    if shape.has_text_frame and shape.text_frame.text.strip():
        draw_text(page, shape, rect, warnings, slide_no)


def draw_text(page, shape, rect, warnings, slide_no):
    tf = shape.text_frame
    ml = (tf.margin_left or 0) * PT
    mr = (tf.margin_right or 0) * PT
    mt = (tf.margin_top or 0) * PT
    mb = (tf.margin_bottom or 0) * PT
    box = pymupdf.Rect(rect.x0 + ml, rect.y0 + mt, rect.x1 - mr, rect.y1 - mb)
    avail_w = box.width
    lines = []  # (text, size, bold, italic, rgb, align, line_h, extra_space_after)
    for p in tf.paragraphs:
        runs = p.runs
        if not runs:
            continue
        r0 = runs[0]
        size = (r0.font.size.pt if r0.font.size else 18)
        bold = bool(r0.font.bold)
        italic = bool(r0.font.italic)
        try:
            rgb = hex_to_rgb01(str(r0.font.color.rgb))
        except Exception:
            rgb = (0, 0, 0)
        spc = 0
        rPr = r0._r.find(qn("a:rPr"))
        if rPr is not None and rPr.get("spc"):
            spc = int(rPr.get("spc")) / 100
        text = "".join(r.text for r in runs)
        align = {1: "l", 2: "c", 3: "r"}.get(int(p.alignment) if p.alignment is not None else 1, "l")
        ls = p.line_spacing if isinstance(p.line_spacing, float) else 1.0
        line_h = size * LINE_FACTOR * ls
        sa = p.space_after.pt if p.space_after is not None else 0
        _, font = font_for(bold, italic)

        def width(s):
            return font.text_length(s, fontsize=size) * WIDTH_FACTOR + spc * len(s)

        for hard in text.split("\n"):
            words = hard.split(" ")
            cur = ""
            for wd in words:
                trial = (cur + " " + wd) if cur else wd
                if width(trial) <= avail_w or not cur:
                    cur = trial
                else:
                    lines.append((cur, size, bold, italic, rgb, align, line_h, 0))
                    cur = wd
            lines.append((cur, size, bold, italic, rgb, align, line_h, 0))
        lines[-1] = lines[-1][:-1] + (sa,)

    total_h = sum(l[6] + l[7] for l in lines)
    if total_h > box.height + 0.5:
        warnings.append(f"slide {slide_no:02d}: text overflow {total_h - box.height:5.1f} pt in box {box.width:.0f}x{box.height:.0f} pt: "
                        f"\"{lines[0][0][:60]}\"")
    anchor = tf.vertical_anchor
    from pptx.enum.text import MSO_ANCHOR
    if anchor == MSO_ANCHOR.MIDDLE:
        y = box.y0 + (box.height - total_h) / 2
    elif anchor == MSO_ANCHOR.BOTTOM:
        y = box.y1 - total_h
    else:
        y = box.y0
    for txt, size, bold, italic, rgb, align, line_h, sa in lines:
        path, font = font_for(bold, italic)
        tw = font.text_length(txt, fontsize=size) * WIDTH_FACTOR
        if align == "c":
            x = box.x0 + (avail_w - tw) / 2
        elif align == "r":
            x = box.x1 - tw
        else:
            x = box.x0
        baseline = y + size * 0.95
        pt = pymupdf.Point(x, baseline)
        page.insert_text(pt, txt, fontsize=size, fontfile=path, fontname="F" + os.path.basename(path)[:-4],
                         color=rgb, morph=(pt, pymupdf.Matrix(WIDTH_FACTOR, 1)))
        y += line_h + sa


def render(pptx_path, out_dir, dpi):
    prs = Presentation(pptx_path)
    W, H = prs.slide_width * PT, prs.slide_height * PT
    doc = pymupdf.open()
    warnings = []
    os.makedirs(out_dir, exist_ok=True)
    for i, slide in enumerate(prs.slides, 1):
        page = doc.new_page(width=W, height=H)
        bg = slide.background.fill
        try:
            page.draw_rect(page.rect, color=None, fill=hex_to_rgb01(str(bg.fore_color.rgb)))
        except Exception:
            page.draw_rect(page.rect, color=None, fill=(1, 1, 1))
        for shape in slide.shapes:
            try:
                draw_shape(page, shape, warnings, i)
            except Exception as e:  # keep going, report
                warnings.append(f"slide {i:02d}: could not draw {shape.shape_type} '{shape.name}': {e}")
        page.get_pixmap(dpi=dpi).save(os.path.join(out_dir, f"slide_{i:02d}.png"))
    doc.save(os.path.join(out_dir, "preview.pdf"))
    return warnings


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("pptx")
    ap.add_argument("--out", default="preview")
    ap.add_argument("--dpi", type=int, default=96)
    a = ap.parse_args()
    ws = render(a.pptx, a.out, a.dpi)
    print(f"rendered to {a.out}/ ({'no' if not ws else len(ws)} warnings)")
    for w in ws:
        print("  !", w)
