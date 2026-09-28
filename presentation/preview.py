"""
Approximate PPTX renderer used to visually QA the deck in an environment with
no PowerPoint/LibreOffice.  It reproduces backgrounds, autoshapes, freeform
charts, text (with wrapping) and tables closely enough to catch overflow,
overlap and spacing problems.  DejaVu Sans is wider than Segoe UI, so anything
that fits here fits in PowerPoint.

Usage:  python3 presentation/preview.py <deck.pptx> <outdir> [slide numbers...]
"""
from __future__ import annotations

import sys
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Emu
from pptx.oxml.ns import qn

DPI = 100
FONT_REG = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
FONT_BOLD = "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"
_font_cache: dict = {}

WARN = (211, 43, 54)


def F(size_pt: float, bold: bool = False) -> ImageFont.FreeTypeFont:
    px = max(7, int(round(size_pt * DPI / 72.0)))
    key = (px, bold)
    if key not in _font_cache:
        _font_cache[key] = ImageFont.truetype(FONT_BOLD if bold else FONT_REG, px)
    return _font_cache[key]


def emu2px(v) -> float:
    return float(v) / 914400.0 * DPI


def color_of(fill):
    try:
        if fill.type is None:
            return None
        from pptx.enum.dml import MSO_FILL
        if fill.type == MSO_FILL.SOLID:
            c = fill.fore_color
            try:
                return "#" + str(c.rgb)
            except Exception:
                return None
    except Exception:
        return None
    return None


def line_of(shape):
    try:
        ln = shape.line
        from pptx.enum.dml import MSO_FILL
        if ln.fill.type == MSO_FILL.SOLID:
            try:
                return "#" + str(ln.color.rgb), max(1, int(round((ln.width or Emu(9525)) / 12700.0 * DPI / 72.0)))
            except Exception:
                return None, 0
    except Exception:
        pass
    return None, 0


def wrap(text, font, maxw):
    words = text.split(" ")
    lines, cur = [], ""
    for w in words:
        trial = w if not cur else cur + " " + w
        if font.getlength(trial) <= maxw or not cur:
            cur = trial
        else:
            lines.append(cur)
            cur = w
    lines.append(cur)
    return lines


def draw_text(dr: ImageDraw.ImageDraw, shape, box, problems, idx):
    tf = shape.text_frame
    x0, y0, x1, y1 = box
    wpx = x1 - x0
    hpx = y1 - y0
    anchor = tf._txBody.find(qn("a:bodyPr")).get("anchor", "t")
    # gather paragraphs
    paras = []
    for p in tf.paragraphs:
        runs = [r for r in p.runs if r.text != ""]
        if not runs:
            paras.append(([], p))
            continue
        paras.append((runs, p))
    total = 0
    rendered = []
    for runs, p in paras:
        size = max([(r.font.size.pt if r.font.size else 12) for r in runs] or [12])
        line_sp = p.line_spacing or 1.0
        if isinstance(line_sp, float):
            lh = size * line_sp * 1.20
        else:
            lh = size * 1.20
        font = F(size, any(r.font.bold for r in runs))
        txt = "".join(r.text for r in runs)
        lines = wrap(txt, font, wpx + 1) if txt else [""]
        rendered.append((runs, p, size, lh, lines, font))
        total += lh * len(lines) + (p.space_after.pt if p.space_after else 0)
    if anchor == "ctr":
        cy = y0 + (hpx - total) / 2.0
    elif anchor == "b":
        cy = y0 + hpx - total
    else:
        cy = y0
    overflow = total > hpx + 1.0
    for runs, p, size, lh, lines, font in rendered:
        align = str(p.alignment) if p.alignment is not None else "LEFT"
        for ln in lines:
            lw = font.getlength(ln)
            if "CENTER" in align:
                cx = x0 + (wpx - lw) / 2.0
            elif "RIGHT" in align:
                cx = x1 - lw
            else:
                cx = x0
            col = color_of(runs[0].font.color) if runs and hasattr(runs[0].font, "color") else None
            try:
                col = "#" + str(runs[0].font.color.rgb)
            except Exception:
                col = col or "#000000"
            dr.text((cx, cy), ln, font=font, fill=col)
            cy += lh
    if overflow:
        problems.append(f"slide{idx+1}: text overflow (+{int(total-hpx)}px) :: {tf.text[:60]!r}")
        dr.rectangle([x0, y0, x1, y1], outline=WARN, width=1)


def custgeom_points(shape):
    spPr = shape._element.spPr
    cg = spPr.find(qn("a:custGeom"))
    if cg is None:
        return None
    path = cg.find(qn("a:pathLst")).find(qn("a:path"))
    pw = float(path.get("w") or 0) or 1.0
    ph = float(path.get("h") or 0) or 1.0
    sw, sh = float(shape.width), float(shape.height)
    sx, sy = float(shape.left), float(shape.top)
    pts = []
    for el in path:
        tag = el.tag.split("}")[-1]
        if tag in ("moveTo", "lnTo"):
            pt = el.find(qn("a:pt"))
            if pt is None:
                continue
            px = float(pt.get("x")) / pw * sw + sx
            py = float(pt.get("y")) / ph * sh + sy
            pts.append((emu2px(px), emu2px(py)))
    return pts


def render(pptx_path: str, outdir: str, only=None):
    prs = Presentation(pptx_path)
    out = Path(outdir)
    out.mkdir(parents=True, exist_ok=True)
    problems = []
    for i, slide in enumerate(prs.slides):
        if only and (i + 1) not in only:
            continue
        img = Image.new("RGB", (int(W_emu(prs, "w")), int(W_emu(prs, "h"))), "white")
        dr = ImageDraw.Draw(img)
        # slide background
        bg = slide.background
        bc = color_of(bg.fill) if bg is not None else None
        if bc:
            dr.rectangle([0, 0, img.width, img.height], fill=bc)
        for shape in slide.shapes:
            st = shape.shape_type
            try:
                box = (emu2px(shape.left), emu2px(shape.top),
                       emu2px(shape.left + shape.width), emu2px(shape.top + shape.height))
            except Exception:
                continue
            if box[0] > img.width + 2 or box[1] > img.height + 2 or box[2] < -2 or box[3] < -2:
                problems.append(f"slide{i+1}: shape off-canvas {shape.shape_type} {shape.name}")
            if str(st) == "TABLE (19)" and shape.has_table:
                tbl = shape.table
                cols = [emu2px(c.width) for c in tbl.columns]
                rows = [emu2px(r.height) for r in tbl.rows]
                cy = box[1]
                for ri, rh in enumerate(rows):
                    cx = box[0]
                    for ci, cw in enumerate(cols):
                        cell = tbl.cell(ri, ci)
                        cf = color_of(cell.fill)
                        if cf:
                            dr.rectangle([cx, cy, cx + cw, cy + rh], fill=cf)
                        txt = cell.text_frame.text
                        sz = 10
                        bold = False
                        col = "#1A1440"
                        for p in cell.text_frame.paragraphs:
                            for r in p.runs:
                                sz = r.font.size.pt if r.font.size else 10
                                bold = bool(r.font.bold)
                                try:
                                    col = "#" + str(r.font.color.rgb)
                                except Exception:
                                    col = "#3D3566"
                        f = F(sz, bold)
                        lines = wrap(txt, f, cw - emu2px(cell.margin_left + cell.margin_right))
                        lh = sz * 1.2 * DPI / 72.0
                        ty = cy + (rh - lh * len(lines)) / 2
                        for ln in lines:
                            dr.text((cx + emu2px(cell.margin_left), ty), ln, font=f, fill=col)
                            ty += lh
                        dr.line([cx, cy, cx, cy + rh], fill="#EFEEF7", width=1)
                        cx += cw
                    dr.line([box[0], cy + rh, box[2], cy + rh], fill="#D8D4EC", width=1)
                    cy += rh
                continue
            name = str(st)
            pts = custgeom_points(shape) if "FREEFORM" in name else None
            if pts:
                col, lw = line_of(shape)
                if len(pts) > 1:
                    dr.line(pts, fill=col or "#000000", width=max(1, lw))
                continue
            if "LINE" in name or "CONNECTOR" in name:
                col, lw = line_of(shape)
                if col:
                    dr.line([box[0], box[1], box[2], box[3]], fill=col, width=max(1, lw))
                continue
            is_oval = False
            try:
                from pptx.enum.shapes import MSO_SHAPE
                is_oval = shape.auto_shape_type == MSO_SHAPE.OVAL
            except Exception:
                pass
            fc = color_of(shape.fill)
            lc, lw = line_of(shape)
            if fc or lc:
                if is_oval:
                    dr.ellipse(box, fill=fc, outline=lc, width=max(1, lw) if lc else 0)
                else:
                    r = 8 if ("ROUNDED" in name) else 0
                    dr.rounded_rectangle(box, radius=r, fill=fc,
                                         outline=lc, width=max(1, lw) if lc else 0)
            if shape.has_text_frame and shape.text_frame.text.strip():
                draw_text(dr, shape, box, problems, i)
        img.save(out / f"slide{i+1:02d}.png")
    for p in problems:
        print(p)
    print(f"rendered {len(only) if only else len(prs.slides.__iter__.__self__._sldIdLst)} pages -> {out}")


def W_emu(prs, which):
    return (emu2px(prs.slide_width) if which == "w" else emu2px(prs.slide_height))


if __name__ == "__main__":
    deck = sys.argv[1]
    od = sys.argv[2]
    nums = [int(a) for a in sys.argv[3:]] or None
    render(deck, od, nums)
