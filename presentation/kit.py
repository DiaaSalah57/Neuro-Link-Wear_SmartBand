"""
NeuroLink Wear — presentation design kit.

Minimalist deck system derived from the dashboard design tokens
(static/css/styles.css) so the deck and the product share one visual language.

Everything produced here is native PowerPoint: real shapes, real text boxes,
real tables — no flattened images — so the file stays fully editable.
"""
from __future__ import annotations

from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.oxml.ns import qn, nsdecls
from pptx.oxml import parse_xml

# ─────────────────────────────── canvas ────────────────────────────────
W, H = 13.333, 7.5
ML, MR = 0.78, 0.78
CW = W - ML - MR                      # 11.773"
COL = CW / 12.0                       # grid column
HEAD_TOP = 1.80                       # first content row when a subtitle exists
FOOT_Y = 7.03
FONT = "Segoe UI"

# ────────────────────── dashboard palette (light) ──────────────────────
BG = "F2F1FA"          # --bg
SURF = "FFFFFF"        # --surface
SURF2 = "F7F6FC"       # --surface-2
SURF3 = "ECEAF6"       # --surface-3
BORDER = "E3E0F2"      # --border
BORDER_STRONG = "CFC9E8"
INK = "1A1440"         # --text
TEXT2 = "3D3566"       # --text-2
MUTED = "7E76A0"       # --muted
ACCENT = "241483"      # --accent
ACCENT2 = "170C5E"     # --accent-2
ACCENT3 = "4C3BD4"
PURPLE = "6B4FD0"      # --purple
LAV = "9D93FF"
LAV_SOFT = "EAE6FB"
OK = "1E8A5F"
OK_BG = "E8F4EF"
WARN = "C97A14"
WARN_BG = "FBF1E2"
DANGER = "D32B36"
DANGER_BG = "FBEAEC"
INFO = "4A3FC0"
INFO_BG = "EDEBFA"

# ───────────────────────── dark canvases ───────────────────────────────
DARK = "0D0A24"
DARK2 = "171240"
DARK3 = "1D1750"
DARK_BORDER = "2E2578"
DARK_GHOST = "241B5C"
LIGHT = "EFEDFB"
LIGHT2 = "C5BFE8"
DARK_MUTED = "8E86B8"

ALIGN = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT, "j": PP_ALIGN.JUSTIFY}
ANCHOR = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}


def gx(i: float) -> float:
    """Grid x for column i (0-based, fractional allowed)."""
    return ML + i * COL


def gw(n: float, gap: float = 0.0) -> float:
    """Width spanning n columns."""
    return n * COL - gap


# ───────────────────────────── primitives ──────────────────────────────
def rgb(h: str) -> RGBColor:
    return RGBColor.from_string(h)


def slide(prs, background: str = BG):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s, background)
    return s


def bg(s, color: str):
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(color)


def rect(s, x, y, w, h, fill=None, line=None, lw=0.75, radius=True,
         adj=0.06, dash=None, transparency=None, kind=None):
    st = kind or (MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE)
    sh = s.shapes.add_shape(st, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.shadow.inherit = False
    if fill:
        sh.fill.solid()
        sh.fill.fore_color.rgb = rgb(fill)
        if transparency:
            sh.fill.transparency = transparency
    else:
        sh.fill.background()
    if line:
        sh.line.color.rgb = rgb(line)
        sh.line.width = Pt(lw)
        if dash:
            sh.line.dash_style = dash
    else:
        sh.line.fill.background()
    if st == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            sh.adjustments[0] = adj
        except Exception:
            pass
    sh.text_frame.word_wrap = True
    return sh


def oval(s, x, y, w, h, fill, line=None, lw=0.75):
    sh = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.shadow.inherit = False
    sh.fill.solid()
    sh.fill.fore_color.rgb = rgb(fill)
    if line:
        sh.line.color.rgb = rgb(line)
        sh.line.width = Pt(lw)
    else:
        sh.line.fill.background()
    return sh


def hline(s, x, y, w, color=BORDER, lw=0.75, dash=None):
    ln = s.shapes.add_connector(1, Inches(x), Inches(y), Inches(x + w), Inches(y))
    ln.line.color.rgb = rgb(color)
    ln.line.width = Pt(lw)
    if dash:
        ln.line.dash_style = dash
    return ln


def vline(s, x, y, h, color=BORDER, lw=0.75, dash=None):
    ln = s.shapes.add_connector(1, Inches(x), Inches(y), Inches(x), Inches(y + h))
    ln.line.color.rgb = rgb(color)
    ln.line.width = Pt(lw)
    if dash:
        ln.line.dash_style = dash
    return ln


def _set_track(run, track):
    if track:
        run.font._rPr.set("spc", str(int(track * 100)))


def tx(s, x, y, w, h, text, size=11, color=INK, bold=False, italic=False,
       align="l", anchor="t", line=1.22, track=0, font=FONT, wrap=True,
       space_after=0, bullet=False, dot="—", dot_color=None):
    """Single-paragraph (or \n separated) text box."""
    box = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = ANCHOR[anchor]
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    lines = text.split("\n") if isinstance(text, str) else list(text)
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = ALIGN[align]
        p.line_spacing = line
        if space_after:
            p.space_after = Pt(space_after)
        if bullet:
            pPr = p._pPr if p._pPr is not None else p._p.get_or_add_pPr()
            pPr.set("marL", "171450")
            pPr.set("indent", "-171450")
            r = p.add_run()
            r.text = dot + "  "
            r.font.size = Pt(size)
            r.font.bold = False
            r.font.name = font
            r.font.color.rgb = rgb(dot_color or color)
            _set_track(r, track)
        r = p.add_run()
        r.text = ln
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.italic = italic
        r.font.name = font
        r.font.color.rgb = rgb(color)
        _set_track(r, track)
    return box


def rich(s, x, y, w, h, paras, anchor="t", line=1.22):
    """
    Multi-run rich text.
    paras: list of paragraph dicts {runs:[{t,size,bold,color,italic,track}],
                                   align, line, space_before, space_after, marL, indent}
    """
    box = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = ANCHOR[anchor]
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = ALIGN[para.get("align", "l")]
        p.line_spacing = para.get("line", line)
        if para.get("space_before"):
            p.space_before = Pt(para["space_before"])
        if para.get("space_after"):
            p.space_after = Pt(para["space_after"])
        if para.get("marL"):
            pPr = p._p.get_or_add_pPr()
            pPr.set("marL", str(int(para["marL"] * 914400)))
            pPr.set("indent", str(int(para.get("indent", -para["marL"]) * 914400)))
        for rn in para.get("runs", []):
            r = p.add_run()
            r.text = rn.get("t", "")
            r.font.size = Pt(rn.get("size", 10))
            r.font.bold = rn.get("bold", False)
            r.font.italic = rn.get("italic", False)
            r.font.name = FONT
            r.font.color.rgb = rgb(rn.get("color", TEXT2))
            _set_track(r, rn.get("track", 0))
    return box


def bullets(s, x, y, w, h, items, size=10.5, color=TEXT2, dot="—",
            dot_color=ACCENT, gap=6, line=1.32, lead_color=INK):
    """Bulleted list; items may be str or (bold_lead, rest)."""
    paras = []
    for i, it in enumerate(items):
        runs = []
        if isinstance(it, tuple):
            lead, rest = it
            runs.append({"t": lead, "size": size, "bold": True, "color": lead_color})
            if rest:
                runs.append({"t": rest, "size": size, "color": color})
        else:
            runs.append({"t": it, "size": size, "color": color})
        paras.append({"runs": runs, "marL": 0.22, "indent": -0.22,
                      "line": line, "space_after": gap if i < len(items) - 1 else 0})
    # dot run injected separately to control colour
    box = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = para["line"]
        if para["space_after"]:
            p.space_after = Pt(para["space_after"])
        pPr = p._p.get_or_add_pPr()
        pPr.set("marL", str(int(0.24 * 914400)))
        pPr.set("indent", str(int(-0.24 * 914400)))
        d = p.add_run()
        d.text = dot + "   "
        d.font.size = Pt(size)
        d.font.name = FONT
        d.font.color.rgb = rgb(dot_color)
        for rn in para["runs"]:
            r = p.add_run()
            r.text = rn["t"]
            r.font.size = Pt(rn["size"])
            r.font.bold = rn.get("bold", False)
            r.font.name = FONT
            r.font.color.rgb = rgb(rn["color"])
    return box


def card(s, x, y, w, h, fill=SURF, line=BORDER, lw=0.75, adj=0.06):
    return rect(s, x, y, w, h, fill=fill, line=line, lw=lw, adj=adj)


def chip(s, x, y, w, h, text, fg=ACCENT, bgc=LAV_SOFT, size=8.5, bold=True,
         track=30, line=None, align="c"):
    rect(s, x, y, w, h, fill=bgc, line=line, adj=0.5)
    return tx(s, x, y + 0.012, w, h - 0.02, text, size=size, color=fg, bold=bold,
              align=align, anchor="m", track=track)


def dot(s, x, y, d, color):
    return oval(s, x, y, d, d, color)


# ───────────────────────────── headings ────────────────────────────────
def header(s, eyebrow, title, subtitle=None, sec=None, num=None, dark=False,
           title_size=26):
    """Standard slide header. Returns y where content may start."""
    tcol = LIGHT if dark else INK
    scol = LIGHT2 if dark else TEXT2
    ecol = LAV if dark else PURPLE
    rule = DARK_BORDER if dark else LAV
    # eyebrow + accent tick
    rect(s, ML, 0.40, 0.30, 0.055, fill=ACCENT3 if not dark else LAV, radius=True, adj=0.5)
    tx(s, ML + 0.42, 0.30, 9.5, 0.26, eyebrow.upper(), size=9, color=ecol, bold=True, track=140)
    tx(s, ML, 0.60, CW - 0.2, 0.62, title, size=title_size, color=tcol, bold=True, line=1.06)
    if subtitle:
        tx(s, ML, 1.26, CW - 1.1, 0.42, subtitle, size=11.5, color=scol, line=1.25)
    if num is not None:
        tx(s, W - MR - 2.2, 0.30, 2.2, 0.26, f"{num:02d}", size=9, color=MUTED if not dark else DARK_MUTED,
           align="r", track=140, bold=True)
    return HEAD_TOP if subtitle else 1.36


def footer(s, num, label="NEUROLINK WEAR", right=None, dark=False):
    col = DARK_BORDER if dark else BORDER
    hline(s, ML, FOOT_Y, CW, color=col, lw=0.75)
    tx(s, ML, FOOT_Y + 0.10, 7.0, 0.24, label.upper(), size=7.5,
       color=DARK_MUTED if dark else MUTED, track=110)
    tx(s, W - MR - 5.0, FOOT_Y + 0.10, 4.6, 0.24,
       (right or "").upper(), size=7.5, color=DARK_MUTED if dark else MUTED,
       align="r", track=110)
    tx(s, W - MR - 0.5, FOOT_Y + 0.10, 0.5, 0.24, f"{num:02d}", size=7.5,
       color=DARK_MUTED if dark else MUTED, align="r", bold=True)


def notes(s, text):
    s.notes_slide.notes_text_frame.text = text


# ───────────────────────────── data viz ────────────────────────────────
def spark(s, x, y, w, h, values, color=ACCENT, lw=1.4, fill=None, dots=False,
          smooth=True, area=False):
    """Tiny native freeform line chart."""
    if not values:
        return None
    lo, hi = min(values), max(values)
    rng = (hi - lo) or 1.0
    n = len(values)
    pts = []
    for i, v in enumerate(values):
        px = x + (w * i / (n - 1) if n > 1 else 0)
        py = y + h - (h * (v - lo) / rng)
        pts.append((px, py))
    if smooth and len(pts) > 3:
        sp = pts[:1]
        for i in range(len(pts) - 1):
            x0, y0 = pts[i]
            x1, y1 = pts[i + 1]
            sp.append(((x0 + x1) / 2, (y0 + y1) / 2))
            sp.append((x1, y1))
        pts = sp
    # NOTE: freeform local units are EMU here (scale=1.0) — python-pptx rounds the
    # start point to an integer, so inch-based local units would snap to whole inches.
    emu = [(int(round(px * 914400)), int(round(py * 914400))) for px, py in pts]
    b = s.shapes.build_freeform(emu[0][0], emu[0][1], scale=1.0)
    b.add_line_segments(emu[1:], close=False)
    sh = b.convert_to_shape()
    sh.shadow.inherit = False
    sh.fill.background()
    sh.line.color.rgb = rgb(color)
    sh.line.width = Pt(lw)
    return sh


def bars(s, x, y, w, h, values, labels=None, colors=None, maxv=None,
         value_labels=None, label_size=8.5, gap=0.16, radius=True, cap=None):
    """Vertical bar group built from native rounded rectangles."""
    n = len(values)
    maxv = maxv or (max(values) * 1.18 or 1)
    bw = (w - gap * (n - 1)) / n
    for i, v in enumerate(values):
        bh = max(0.035, h * (v / maxv))
        bx = x + i * (bw + gap)
        by = y + h - bh
        c = (colors[i] if isinstance(colors, (list, tuple)) else colors) or ACCENT
        rect(s, bx, by, bw, bh, fill=c, radius=radius, adj=0.18)
        if value_labels:
            tx(s, bx, by - 0.26, bw, 0.22, value_labels[i], size=9, color=INK,
               bold=True, align="c")
        if labels:
            tx(s, bx, y + h + 0.08, bw, 0.22, labels[i], size=label_size,
               color=MUTED, align="c", track=40)
    if cap:
        tx(s, x, y - 0.26, w, 0.22, cap, size=8, color=MUTED, align="r", track=60)


def hbar(s, x, y, w, h, pct, color=ACCENT, bgc=SURF3, adj=0.5):
    rect(s, x, y, w, h, fill=bgc, radius=True, adj=adj)
    if pct > 0:
        rect(s, x, y, max(0.05, w * min(pct, 1.0)), h, fill=color, radius=True, adj=adj)


def progress_row(s, x, y, w, label, value_txt, pct, color=ACCENT, size=10,
                 label_w=None):
    lw = label_w or w * 0.44
    tx(s, x, y, lw, 0.24, label, size=size, color=TEXT2, anchor="m")
    bar_x = x + lw + 0.16
    bar_w = w - lw - 0.16 - 0.96
    hbar(s, bar_x, y + 0.055, bar_w, 0.13, pct, color=color)
    tx(s, x + w - 0.92, y, 0.92, 0.30, value_txt, size=size, color=INK, bold=True,
       align="r", anchor="m")


# ───────────────────────────── tables ──────────────────────────────────
_ORDER = ["a:lnL", "a:lnR", "a:lnT", "a:lnB", "a:lnTlToBr", "a:lnBlToTr",
          "a:cell3D", "a:noFill", "a:solidFill", "a:gradFill", "a:blipFill",
          "a:pattFill", "a:grpFill", "a:headers", "a:extLst"]


def _cell_border(cell, edge, color=BORDER, lw=0.75):
    tcPr = cell._tc.get_or_add_tcPr()
    name = f"a:ln{edge}"
    old = tcPr.find(qn(name))
    if old is not None:
        tcPr.remove(old)
    el = parse_xml(
        f'<{name} {nsdecls("a")} w="{int(lw * 12700)}" cap="flat" cmpd="sng" algn="ctr">'
        f'<a:solidFill><a:srgbClr val="{color}"/></a:solidFill>'
        f'<a:prstDash val="solid"/></{name}>')
    idx = _ORDER.index(f"a:ln{edge}")
    pos = len(tcPr)
    for i, child in enumerate(tcPr):
        ctag = child.tag.split("}")[-1]
        cname = f"a:{ctag}"
        if cname in _ORDER and _ORDER.index(cname) > idx:
            pos = i
            break
    tcPr.insert(pos, el)


def table(s, x, y, col_w, rows, row_h=0.30, head_h=0.32, size=9.5, head_size=8.5,
          header=True, align=None, head_align=None, bold_col0=False, zebra=None,
          row_colors=None, pad=0.10, valign="m", line=BORDER, head_line=BORDER_STRONG,
          fill=None, head_fill=None, wrap_head=True):
    """
    Minimalist native table: hairline bottom rules, muted uppercase header.
    col_w / rows are in inches and text values.
    """
    n_rows, n_cols = len(rows), len(col_w)
    tw = sum(col_w)
    th = (head_h if header else 0) + row_h * (n_rows - (1 if header else 0))
    gf = s.shapes.add_table(n_rows, n_cols, Inches(x), Inches(y), Inches(tw), Inches(th))
    tbl = gf.table
    tbl.first_row = False
    tbl.horz_banding = False
    # kill inherited table style so our fills show cleanly
    tblPr = tbl._tbl.find(qn("a:tblPr"))
    if tblPr is not None:
        for st in tblPr.findall(qn("a:tableStyleId")):
            tblPr.remove(st)
        tblPr.set("firstRow", "0")
        tblPr.set("bandRow", "0")
    for i, cw in enumerate(col_w):
        tbl.columns[i].width = Inches(cw)
    for r in range(n_rows):
        tbl.rows[r].height = Inches(head_h if (header and r == 0) else row_h)
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = tbl.cell(r, c)
            cell.margin_left = cell.margin_right = Inches(pad)
            cell.margin_top = cell.margin_bottom = Inches(0.02)
            cell.vertical_anchor = ANCHOR[valign]
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.line_spacing = 1.06
            is_head = header and r == 0
            al = (head_align if is_head else align)
            al = al[c] if isinstance(al, (list, tuple)) else (al or "l")
            p.alignment = ALIGN[al]
            run = p.add_run()
            run.text = str(val)
            run.font.size = Pt(head_size if is_head else size)
            run.font.bold = is_head or (bold_col0 and c == 0)
            run.font.name = FONT
            run.font.color.rgb = rgb(MUTED if is_head else (INK if (bold_col0 and c == 0) else TEXT2))
            if is_head:
                _set_track(run, 60)
            # fills
            wanted = head_fill if is_head else None
            if fill or wanted or zebra or row_colors:
                cf = wanted or (row_colors[r] if row_colors else None) or \
                     (zebra if (zebra and r % 2 == 0) else fill)
                if cf:
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = rgb(cf)
                else:
                    cell.fill.background()
            else:
                cell.fill.background()
            # hairlines
            bottom = head_line if is_head else line
            _cell_border(cell, "B", bottom, 0.75 if is_head else 0.5)
            if c > 0:
                _cell_border(cell, "L", line, 0.4)
    return tbl


# ───────────────────────── slide furniture ─────────────────────────────
def divider(prs, slide_no, section, title, subtitle, items=None, label=None):
    """Full-bleed dark section divider."""
    s = slide(prs, DARK)
    rect(s, 0, 0, W, H, fill=DARK, radius=False)
    rect(s, 0, 0, 0.10, H, fill=ACCENT3, radius=False)
    tx(s, ML + 0.30, 0.95, 4.0, 1.9, str(section), size=112, color=DARK_GHOST, bold=True, line=0.9)
    tx(s, ML + 0.34, 3.05, 7.6, 0.30, (label or "SECTION").upper(), size=9.5,
       color=LAV, bold=True, track=180)
    tx(s, ML + 0.30, 3.38, 6.55, 0.9, title, size=32, color=LIGHT, bold=True, line=1.06)
    tx(s, ML + 0.34, 4.40, 6.1, 0.7, subtitle, size=12, color=LIGHT2, line=1.35)
    if items:
        iy = 2.42
        for it in items:
            tx(s, W - MR - 4.55, iy, 4.55, 0.30, it, size=10, color=LIGHT2, anchor="m")
            hline(s, W - MR - 4.55, iy + 0.46, 4.55, color=DARK_BORDER, lw=0.75)
            iy += 0.70
    footer(s, slide_no, "NEUROLINK WEAR", right="PROTOTYPE &nbsp;·&nbsp; DISCUSSION MATERIAL".replace("&nbsp;", " "), dark=True)
    return s


def stat_card(s, x, y, w, h, value, label, sub=None, color=ACCENT, accent_bar=None):
    card(s, x, y, w, h)
    if accent_bar:
        rect(s, x, y, 0.045, h, fill=accent_bar, radius=True, adj=0.5)
    tx(s, x + 0.24, y + 0.20, w - 0.4, 0.5, value, size=25, color=color, bold=True, line=1.0)
    tx(s, x + 0.24, y + 0.72, w - 0.4, 0.52, label, size=10, color=TEXT2, line=1.22)
    if sub:
        tx(s, x + 0.24, y + h - 0.42, w - 0.4, 0.3, sub, size=8.5, color=MUTED, track=20)


def kv_row(s, x, y, w, key, val, key_w=2.4, size=10, val_color=INK, key_color=MUTED):
    tx(s, x, y, key_w, 0.28, key, size=size, color=key_color, anchor="m")
    tx(s, x + key_w, y, w - key_w, 0.28, val, size=size, color=val_color, anchor="m")


def badge(s, x, y, w, h, text, color=DANGER, bgc=DANGER_BG, size=8.5):
    chip(s, x, y, w, h, text, fg=color, bgc=bgc, size=size, track=40)
