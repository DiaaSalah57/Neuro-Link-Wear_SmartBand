from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls
from pathlib import Path

# NeuroLink Wear investor / technology pitch deck.
# Built from native PowerPoint shapes and text for straightforward editing.
OUT = Path(__file__).resolve().parents[1] / 'NeuroLink_Wear_Investor_Deck.pptx'

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]

# Dashboard-derived palette (static/css/styles.css)
C = {
    'bg': 'F2F1FA', 'surface': 'FFFFFF', 'surface2': 'F7F6FC', 'surface3': 'ECEAF6',
    'border': 'E3E0F2', 'ink': '1A1440', 'text2': '3D3566', 'muted': '7E76A0',
    'purple': '241483', 'purple2': '170C5E', 'accent': '4C3BD4', 'lavender': '6B4FD0',
    'lavender2': '9D93FF', 'green': '1E8A5F', 'greenbg': 'E8F4EF',
    'amber': 'C97A14', 'amberbg': 'FBF1E2', 'red': 'D32B36', 'redbg': 'FBEAEC',
    'dark': '0D0A24', 'darkcard': '171240', 'darkborder': '2E2578', 'white': 'FFFFFF',
    'gray': 'A8A2BE', 'blue': '4A3FC0'
}
W, H = 13.333, 7.5
FONT = 'Aptos'


def rgb(hexv):
    return RGBColor.from_string(hexv)


def set_bg(slide, color):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(color)


def shape(slide, x, y, w, h, fill=None, line=None, radius=True, transparency=0, kind=None, line_width=1):
    st = kind or (MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE)
    sh = slide.shapes.add_shape(st, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill:
        sh.fill.solid(); sh.fill.fore_color.rgb = rgb(fill)
        if transparency: sh.fill.transparency = transparency
    else:
        sh.fill.background()
    if line:
        sh.line.color.rgb = rgb(line); sh.line.width = Pt(line_width)
    else:
        sh.line.fill.background()
    if st == MSO_SHAPE.ROUNDED_RECTANGLE:
        try: sh.adjustments[0] = 0.12
        except Exception: pass
    return sh


def text(slide, x, y, w, h, txt, size=14, color=None, bold=False, align=PP_ALIGN.LEFT,
         valign=MSO_ANCHOR.TOP, margin=0, font=FONT, italic=False, spacing=1.0):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.clear(); tf.word_wrap = True
    tf.margin_left = Inches(margin); tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin); tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]; p.alignment = align; p.space_after = Pt(0); p.space_before = Pt(0)
    p.line_spacing = spacing
    r = p.add_run(); r.text = txt
    r.font.name = font; r.font.size = Pt(size); r.font.bold = bold; r.font.italic = italic
    r.font.color.rgb = rgb(color or C['ink'])
    return tb


def line(slide, x1, y1, x2, y2, color=None, width=1.2, dash=None, begin=None, end=None):
    ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    ln.line.color.rgb = rgb(color or C['border']); ln.line.width = Pt(width)
    if dash: ln.line.dash_style = dash
    if end:
        ln.line.end_arrowhead = end
    if begin:
        ln.line.begin_arrowhead = begin
    return ln


def pill(slide, x, y, w, label, fill=None, color=None, h=0.30, size=8.5, outline=None):
    sh = shape(slide, x, y, w, h, fill or C['surface3'], outline, True)
    text(slide, x+0.06, y+0.01, w-0.12, h-0.02, label, size, color or C['purple'], True,
         PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    return sh


def card(slide, x, y, w, h, title, body='', accent=None, fill=None, body_size=11.5,
         title_size=14, title_color=None, body_color=None, outline=None):
    shape(slide, x, y, w, h, fill or C['surface'], outline or C['border'])
    if accent:
        shape(slide, x, y+0.16, 0.055, h-0.32, accent, None, False)
    tx = x + (0.22 if not accent else 0.29)
    text(slide, tx, y+0.19, w-(tx-x)-0.20, 0.38, title, title_size, title_color or C['ink'], True)
    if body:
        text(slide, tx, y+0.64, w-(tx-x)-0.20, h-0.79, body, body_size, body_color or C['text2'], False, spacing=1.04)


def bullet_list(slide, x, y, w, items, size=11.5, color=None, gap=0.52, dot_color=None, maxh=None):
    color = color or C['text2']; dot_color = dot_color or C['purple']
    for i, item in enumerate(items):
        yy = y + i*gap
        shape(slide, x, yy+0.10, 0.075, 0.075, dot_color, None, True, kind=MSO_SHAPE.OVAL)
        text(slide, x+0.18, yy, w-0.18, min(gap, maxh or gap), item, size, color, False, spacing=1.02)


def header(slide, section, title_txt, subtitle=None, dark=False, page=1, title_size=27):
    # Keep long titles on one line within the fixed header band.
    n = len(title_txt)
    if n > 84: title_size = min(title_size, 20)
    elif n > 74: title_size = min(title_size, 21)
    elif n > 64: title_size = min(title_size, 23)
    elif n > 54: title_size = min(title_size, 25)
    fg = C['white'] if dark else C['ink']; muted = C['lavender2'] if dark else C['muted']
    shape(slide, 0.52, 0.32, 0.16, 0.16, C['lavender2'] if dark else C['purple'], None, True, kind=MSO_SHAPE.OVAL)
    text(slide, 0.79, 0.25, 2.65, 0.25, 'NEUROLINK  /  WEAR', 9.2, fg, True, valign=MSO_ANCHOR.MIDDLE)
    text(slide, 9.05, 0.26, 3.75, 0.24, section.upper(), 8.5, muted, True, PP_ALIGN.RIGHT, MSO_ANCHOR.MIDDLE)
    text(slide, 0.58, 0.71, 12.1, 0.55, title_txt, title_size, fg, True)
    if subtitle:
        text(slide, 0.60, 1.34, 12.0, 0.35, subtitle, 11.2, muted if dark else C['muted'])
    # footer
    line(slide, 0.58, 7.12, 12.76, 7.12, C['darkborder'] if dark else C['border'], 0.8)
    text(slide, 0.60, 7.20, 4.4, 0.16, 'WORKING PROTOTYPE  •  DISCUSSION MATERIAL', 7.2, muted, True)
    text(slide, 12.10, 7.18, 0.62, 0.18, f'{page:02d}', 8.5, muted, True, PP_ALIGN.RIGHT)


def new_slide(section, title_txt, subtitle=None, page=None, dark=False, title_size=27):
    s = prs.slides.add_slide(blank); set_bg(s, C['dark'] if dark else C['bg'])
    header(s, section, title_txt, subtitle, dark, page or len(prs.slides), title_size)
    return s


def small_label(slide, x, y, w, txt, dark=False):
    text(slide, x, y, w, 0.20, txt.upper(), 8.2, C['lavender2'] if dark else C['muted'], True)


def stat(slide, x, y, w, val, label, accent=None):
    text(slide, x, y, w, 0.55, val, 27, accent or C['purple'], True)
    text(slide, x, y+0.60, w, 0.32, label, 9.5, C['muted'], True)

# 01 — Cover
s = prs.slides.add_slide(blank); set_bg(s, C['dark'])
# decorative glow rings and grid motifs
shape(s, 7.65, 0.55, 4.75, 4.75, C['darkcard'], C['darkborder'], True)
shape(s, 8.08, 0.97, 3.9, 3.9, C['dark'], C['darkborder'], True)
shape(s, 8.47, 1.37, 3.1, 3.1, C['darkcard'], None, True)
shape(s, 8.82, 1.72, 2.4, 2.4, C['dark'], C['darkborder'], True)
shape(s, 9.49, 2.39, 1.05, 1.05, C['purple'], None, True, kind=MSO_SHAPE.OVAL)
# simple smartband illustration on right
shape(s, 9.67, 0.72, 0.72, 4.48, C['purple2'], C['lavender'], True)
shape(s, 9.20, 2.08, 1.68, 1.64, C['surface'], C['lavender2'], True)
shape(s, 9.36, 2.24, 1.36, 1.32, C['surface2'], C['border'], True)
text(s, 9.49, 2.38, 1.08, 0.20, 'NLW', 10, C['purple'], True, PP_ALIGN.CENTER)
text(s, 9.48, 2.72, 1.10, 0.27, '72', 20, C['ink'], True, PP_ALIGN.CENTER)
text(s, 9.45, 3.11, 1.18, 0.16, 'BPM  •  LIVE', 7, C['muted'], True, PP_ALIGN.CENTER)
# left identity
shape(s, 0.72, 0.62, 0.18, 0.18, C['lavender2'], None, True, kind=MSO_SHAPE.OVAL)
text(s, 1.04, 0.56, 2.8, 0.30, 'NEUROLINK  /  WEAR', 10, C['white'], True)
pill(s, 0.75, 1.55, 2.24, 'INVESTOR + TECHNICAL PITCH', C['darkcard'], C['lavender2'], 0.34, 8.5, C['darkborder'])
text(s, 0.72, 2.25, 6.65, 1.40, 'Care signals.\nClear response.', 36, C['white'], True, spacing=0.93)
text(s, 0.76, 4.02, 5.95, 0.75, 'An explainable wearable health & safety platform\nfor older adults, families and care teams.', 16, 'C5BFE8', False, spacing=1.12)
line(s, 0.77, 5.17, 5.20, 5.17, C['darkborder'], 1)
pill(s, 0.78, 5.48, 1.35, 'ESP32 + SENSORS', C['darkcard'], C['white'], 0.31, 8, C['darkborder'])
pill(s, 2.23, 5.48, 1.40, 'CALIBRATED AI', C['darkcard'], C['white'], 0.31, 8, C['darkborder'])
pill(s, 3.74, 5.48, 1.52, 'SAFETY WORKFLOW', C['darkcard'], C['white'], 0.31, 8, C['darkborder'])
text(s, 0.78, 6.77, 6.5, 0.22, 'Prototype status: software platform + firmware source; clinical validation is future work.', 8.5, C['gray'])
text(s, 11.70, 6.90, 0.90, 0.20, '2026', 9, C['gray'], True, PP_ALIGN.RIGHT)

# 02 — Truthful product snapshot
s = new_slide('01  /  INVESTMENT THESIS', 'A product platform is built. The evidence agenda is clear.',
              'Separate implemented capability from what still needs real-world proof.', page=2)
# three QA metrics
shape(s, 0.62, 1.93, 12.08, 1.08, C['surface'], C['border'])
stat(s, 0.92, 2.12, 2.2, '141 / 141', 'SOFTWARE QA CHECKS PASS', C['purple'])
line(s, 3.42, 2.11, 3.42, 2.86, C['border'], 1)
stat(s, 3.82, 2.12, 2.0, '74 / 74', 'CALIBRATION TESTS PASS', C['lavender'])
line(s, 6.12, 2.11, 6.12, 2.86, C['border'], 1)
stat(s, 6.52, 2.12, 2.0, '17 / 17', 'CALIBRATION UI CHECKS', C['green'])
line(s, 8.82, 2.11, 8.82, 2.86, C['border'], 1)
text(s, 9.18, 2.12, 3.1, 0.30, 'Prototype stage', 16, C['ink'], True)
text(s, 9.18, 2.49, 3.18, 0.32, 'No clinical validation or commercial traction claimed.', 9.5, C['muted'])
card(s, 0.62, 3.28, 5.83, 2.78, 'Built today',
     '• Responsive caregiver dashboard + FastAPI backend\n• Live telemetry simulator, WebSocket stream and event feed\n• ESP32 firmware source + sensor payload ingest path\n• Explainable equation / calibration detection layer\n• Alerts, trends, safety timeline, contacts and device management', C['green'], body_size=11.2)
card(s, 6.68, 3.28, 6.02, 2.78, 'Proof still required',
     '• Sensor accuracy across wearers, motion and skin tones\n• Battery, comfort, adherence and production BOM\n• False-alert rate, sensitivity and escalation latency\n• Clinical, privacy, security and regulatory pathways\n• Live SMS / voice delivery and multi-tenant cloud operations', C['amber'], body_size=11.2)
text(s, 0.64, 6.48, 11.7, 0.30, 'Investor framing: validated software workflow + testable prototype—not a cleared medical device or proven clinical service.', 10.5, C['purple'], True)

# 03 — Problem
s = new_slide('02  /  NEED', 'The gap is not another reading. It is context and response.',
              'Caregivers often operate between scheduled check-ins, fragmented signals and uncertain next steps.', page=3)
problem_cards = [
    ('01', 'Signals are isolated', 'A pulse, oxygen value or motion event is hard to interpret without personal history and activity context.', C['purple']),
    ('02', 'Changes can be subtle', 'Meaningful deviations may emerge as patterns over time—not one threshold crossing.', C['lavender']),
    ('03', 'Alert fatigue is real', 'Generic thresholds can miss personal baselines or create noisy, low-value notifications.', C['amber']),
    ('04', 'Help is a workflow', 'A risk signal only matters if the wearer or trusted contact can act, confirm and follow through.', C['red']),
]
for i,(num,t,b,a) in enumerate(problem_cards):
    x=0.62+(i%2)*6.12; y=2.02+(i//2)*2.12
    shape(s,x,y,5.84,1.82,C['surface'],C['border'])
    text(s,x+0.22,y+0.20,0.62,0.34,num,17,a,True)
    text(s,x+0.95,y+0.18,4.50,0.30,t,15,C['ink'],True)
    text(s,x+0.95,y+0.64,4.48,0.93,b,11.2,C['text2'])
pill(s,0.64,6.46,2.05,'PROBLEM HYPOTHESIS',C['surface3'],C['purple'],0.30,8)
text(s,2.88,6.47,9.65,0.30,'Pilot interviews should quantify the pain, current workarounds and willingness to pay.',10,C['muted'])

# 04 — Audience
s = new_slide('03  /  CUSTOMER', 'Start with the person who needs a safer day—and the person who checks in.',
              'A focused initial use case gives the product a clear buyer, user and success metric.', page=4, title_size=24)
people = [
    ('WEARER', 'Older adult living independently', 'Wants simple, dignified support; comfortable band, clear on-device status and trusted contact options.', C['purple']),
    ('BUYER / USER', 'Family caregiver', 'Needs a calm, at-a-glance view, actionable alerts, shared contacts and less uncertainty between visits.', C['lavender']),
    ('CHANNEL BUYER', 'Home-care or senior-living team', 'Needs a manageable multi-wearer workflow, incident history and escalation ownership.', C['green'])]
for i,(tag,title_,body,ac) in enumerate(people):
    x=0.62+i*4.10
    shape(s,x,2.10,3.83,2.75,C['surface'],C['border'])
    pill(s,x+0.22,2.33,1.28,tag,C['surface3'],ac,0.28,7.8)
    text(s,x+0.22,2.82,3.34,0.74,title_,16,C['ink'],True)
    text(s,x+0.22,3.68,3.30,0.94,body,11.1,C['text2'])
shape(s,0.62,5.10,12.08,1.27,C['purple'],None)
text(s,0.91,5.31,1.88,0.24,'SECONDARY AUDIENCES',8.5,'D6D0FF',True)
text(s,0.91,5.68,11.1,0.45,'Aging-at-home programs  •  Post-discharge support partners*  •  Remote-care providers  •  Wellness / stress management later',12,C['white'],True)
text(s,0.93,6.47,11.5,0.22,'*Only with appropriate clinical partner, evidence and regulatory review. Initial wedge: family caregiving / senior-living pilots.',9,C['muted'])

# 05 — Market framing
s = new_slide('04  /  OPPORTUNITY', 'A bottom-up market case is more credible than an unsupported TAM.',
              'Focus the first market on aging-at-home safety, then expand through care organizations.', page=5)
# left focus ladder
shape(s,0.62,2.02,7.28,4.46,C['surface'],C['border'])
text(s,0.91,2.28,3.5,0.28,'LAND → EXPAND',10,C['purple'],True)
steps=[('01','Family caregiver households','Wearer + 1–3 trusted contacts; direct subscription hypothesis'),
       ('02','Home care & senior living','Multi-wearer dashboard, staff workflow, paid pilots'),
       ('03','Remote care ecosystem','Provider / insurer partnerships after proof and integration')]
for i,(n,t,b) in enumerate(steps):
    yy=2.83+i*1.05
    shape(s,0.95,yy,0.46,0.46,C['surface3'],None,True)
    text(s,1.00,yy+0.09,0.36,0.22,n,9,C['purple'],True,PP_ALIGN.CENTER)
    text(s,1.62,yy,5.55,0.27,t,14,C['ink'],True)
    text(s,1.62,yy+0.34,5.65,0.48,b,10.5,C['text2'])
    if i<2: line(s,1.18,yy+0.49,1.18,yy+0.99,C['border'],1.3)
shape(s,8.22,2.02,4.48,4.46,C['dark'],None)
text(s,8.56,2.34,3.75,0.25,'SCALE ILLUSTRATION',9,C['lavender2'],True)
text(s,8.56,2.83,3.7,0.68,'100,000',30,C['white'],True)
text(s,8.56,3.53,3.7,0.26,'paid wearer accounts',10,'C5BFE8',True)
line(s,8.56,3.99,12.23,3.99,C['darkborder'],1)
text(s,8.56,4.23,3.7,0.66,'$14.4M',27,C['lavender2'],True)
text(s,8.56,4.90,3.6,0.38,'annual subscription revenue at $12/mo',9.5,'C5BFE8')
text(s,8.56,5.57,3.68,0.48,'Illustrative scale case only—not a market-size estimate or forecast.',9,C['gray'])
text(s,0.65,6.69,11.7,0.18,'Next: define launch geography, eligible customer base, channel access, pricing and reimbursement before publishing TAM / SAM.',8.3,C['muted'])

# 06 — Solution
s = new_slide('05  /  SOLUTION', 'One loop: sense, interpret, explain, connect.',
              'NeuroLink Wear combines a sensor stream with personal context and a caregiver action path.', page=6)
flow=[('SENSE','Vitals + motion','PPG · GSR · IMU\nTemperature · GPS',C['purple']),
      ('INTERPRET','Personal context','Baselines · trends\nClinical safety floors',C['lavender']),
      ('EXPLAIN','Human-readable','Why it fired\nWhat to do next',C['amber']),
      ('CONNECT','Care response','Acknowledge · contact\nLocation · event log',C['green'])]
for i,(a,b,c,d) in enumerate(flow):
    x=0.72+i*3.15
    shape(s,x,2.38,2.65,2.30,C['surface'],C['border'])
    shape(s,x+0.22,2.66,0.50,0.50,d,None,True,kind=MSO_SHAPE.OVAL)
    text(s,x+0.22,2.79,0.50,0.20,str(i+1),10,C['white'],True,PP_ALIGN.CENTER)
    text(s,x+0.22,3.37,2.24,0.28,a,9,d,True)
    text(s,x+0.22,3.73,2.30,0.34,b,14,C['ink'],True)
    text(s,x+0.22,4.17,2.30,0.42,c,10,C['text2'])
    if i<3:
        line(s,x+2.68,3.52,x+3.02,3.52,C['lavender'],1.6)
shape(s,0.72,5.15,11.97,1.03,C['surface2'],C['border'])
text(s,0.99,5.39,2.05,0.24,'DESIGN PRINCIPLES',8.3,C['muted'],True)
text(s,3.08,5.32,9.20,0.48,'Explainable over opaque  •  Personal without weakening safety floors  •  Human-led escalation',12,C['purple'],True,PP_ALIGN.CENTER,MSO_ANCHOR.MIDDLE)

# 07 — Product experience
s = new_slide('06  /  PRODUCT', 'Designed for two screens: the wearer’s wrist and the caregiver’s dashboard.',
              'A small on-device status surface paired with a responsive, role-aware web experience.', page=7, title_size=24)
# wearable mock on left
shape(s,0.73,2.00,3.12,4.35,C['dark'],None)
text(s,1.03,2.29,2.53,0.23,'ON-WRIST DISPLAY',8.5,C['lavender2'],True,PP_ALIGN.CENTER)
shape(s,1.51,2.95,1.55,2.45,C['darkcard'],C['lavender'],True)
shape(s,1.68,3.16,1.20,1.98,C['dark'],C['darkborder'],True)
text(s,1.81,3.40,0.93,0.24,'HEART RATE',7,C['gray'],True,PP_ALIGN.CENTER)
text(s,1.82,3.77,0.90,0.45,'72',25,C['white'],True,PP_ALIGN.CENTER)
text(s,1.82,4.28,0.90,0.20,'SpO₂  97%',8,C['lavender2'],True,PP_ALIGN.CENTER)
text(s,1.82,4.63,0.90,0.22,'Resting',8,C['gray'],False,PP_ALIGN.CENTER)
text(s,1.02,5.72,2.50,0.38,'Simple, glanceable status;\nnot a diagnostic display.',9.5,C['white'],False,PP_ALIGN.CENTER)
# dashboard mock right
shape(s,4.18,2.00,8.52,4.35,C['surface'],C['border'])
shape(s,4.18,2.00,8.52,0.45,C['purple'],None,True)
text(s,4.43,2.11,3.10,0.18,'NEUROLINK  /  OVERVIEW',8,C['white'],True)
pill(s,10.44,2.08,1.72,'DEVICE CONNECTED',C['greenbg'],C['green'],0.25,7.1)
# vital cards
for i,(lab,val,unit) in enumerate([('HEART RATE','72','bpm'),('SpO₂','97','%'),('TEMP','36.7','°C'),('STRESS','0.24','index')]):
    x=4.43+i*1.95
    shape(s,x,2.70,1.70,0.96,C['surface2'],C['border'])
    text(s,x+0.12,2.82,1.45,0.15,lab,7,C['muted'],True)
    text(s,x+0.12,3.04,0.86,0.35,val,18,C['ink'],True)
    text(s,x+1.06,3.18,0.48,0.16,unit,7,C['muted'])
# chart and alert preview
shape(s,4.43,3.90,4.30,1.98,C['surface2'],C['border'])
text(s,4.65,4.08,3.70,0.20,'LIVE TELEMETRY  •  TREND',8,C['muted'],True)
for i in range(6):
    line(s,4.67,4.51+i*0.19,8.46,4.51+i*0.19,'E8E5F1',0.5)
pts=[(4.75,5.40),(5.35,5.15),(5.95,5.20),(6.55,4.85),(7.15,4.98),(7.75,4.62),(8.32,4.72)]
for a,b in zip(pts,pts[1:]): line(s,a[0],a[1],b[0],b[1],C['lavender'],2)
shape(s,8.98,3.90,3.44,1.98,C['redbg'],None)
pill(s,9.20,4.10,1.00,'HIGH',C['red'],C['white'],0.25,7.2)
text(s,9.20,4.49,2.94,0.24,'High-stress pattern',12,C['ink'],True)
text(s,9.20,4.83,2.91,0.64,'Personal baseline + elevated GSR and reduced HRV. Check in and confirm wellbeing.',8.7,C['text2'])
text(s,4.43,6.06,7.80,0.18,'Illustrative interface; values are mock data. The app uses seeded demo records in the prototype.',8,C['muted'],False,italic=True)

# 08 — Sensors
s = new_slide('07  /  HARDWARE', 'A multi-sensor stack adds context—but every signal needs validation.',
              'Components referenced by the current ESP32 firmware source; measurements are prototype-grade.', page=8, title_size=24)
sensors=[
 ('PPG MODULE','MAX30105 (firmware)','Heart rate, HRV estimate, SpO₂ estimate','Pulse / oxygen context',C['purple']),
 ('GSR','Analog skin conductance','Conductance + sweat-response proxy','Arousal / stress context',C['lavender']),
 ('IMU','MPU6050','3-axis acceleration + gyroscope','Movement, steps, fall signature',C['green']),
 ('TEMPERATURE','MLX90614 IR sensor','Object / skin-surface temperature','Thermal trend; not core temp',C['amber']),
 ('LOCATION','NEO-6M GPS','Outdoor position fix + timestamp','Context for a safety event',C['red']),
 ('DISPLAY + MCU','ESP32 + SSD1306 OLED','Feature processing, local display, Wi-Fi / MQTT','On-wrist status + data transport',C['blue'])]
for i,(tag,part,signal,use,ac) in enumerate(sensors):
    x=0.62+(i%3)*4.10; y=2.00+(i//3)*1.96
    shape(s,x,y,3.82,1.69,C['surface'],C['border'])
    shape(s,x+0.21,y+0.24,0.42,0.42,ac,None,True,kind=MSO_SHAPE.OVAL)
    text(s,x+0.78,y+0.21,2.82,0.20,tag,8,C['muted'],True)
    text(s,x+0.21,y+0.76,3.40,0.24,part,13,C['ink'],True)
    text(s,x+0.21,y+1.06,3.40,0.20,signal,9.3,C['text2'])
    text(s,x+0.21,y+1.35,3.40,0.18,use,8.4,ac,True)
shape(s,0.62,6.15,12.02,0.57,C['amberbg'],None)
text(s,0.87,6.31,11.4,0.20,'Validation gate: SpO₂ is a simple red/IR estimate today; benchmark HR, HRV, temperature, motion, fit / skin tone, indoor GPS and power before health claims.',8.7,C['amber'],True)

# 09 — Architecture
s = new_slide('08  /  TECHNICAL ARCHITECTURE', 'The end-to-end software path is in place; device timing is a field-readiness gate.',
              'A single ingest and detection pipeline serves both simulator traffic and device payloads.', page=9, title_size=23)
blocks=[('01','ESP32 + sensors','PPG · GSR · IMU\nTemp · GPS',C['purple']),('02','MQTT / ingest','Device payload\nIdentity + parsing',C['lavender']),('03','Calibration','Patient baselines\nSignal context',C['blue']),('04','Detection','Rules + equations\nSeverity / evidence',C['amber']),('05','Store + broadcast','SQLite prototype\nREST + WebSocket',C['green']),('06','Care dashboard','Trends · alerts\nSafety actions',C['red'])]
for i,(n,t,b,ac) in enumerate(blocks):
    x=0.60+i*2.12
    shape(s,x,2.32,1.77,1.77,C['surface'],C['border'])
    pill(s,x+0.17,2.50,0.40,n,C['surface3'],ac,0.25,7.5)
    text(s,x+0.17,2.91,1.45,0.35,t,11.5,C['ink'],True)
    text(s,x+0.17,3.38,1.46,0.50,b,9,C['text2'])
    if i<5: line(s,x+1.79,3.17,x+2.06,3.17,C['lavender'],1.5)
# optional language model branch
shape(s,7.91,4.65,2.39,0.75,C['surface2'],C['border'])
text(s,8.10,4.81,2.03,0.19,'OPTIONAL LLM',8,C['purple'],True,PP_ALIGN.CENTER)
text(s,8.08,5.07,2.07,0.18,'Explanation layer only',8,C['muted'],False,PP_ALIGN.CENTER)
line(s,8.23,4.12,8.23,4.60,C['lavender'],1.2,MSO_LINE_DASH_STYLE.DASH)
shape(s,0.62,4.75,6.78,1.24,C['surface'],C['border'])
text(s,0.87,4.97,2.0,0.22,'DEMO / DEVELOPMENT',8,C['muted'],True)
text(s,0.87,5.34,6.15,0.43,'Simulator → ~2-second dashboard stream • seeded history • scripted fall / stress / fever / low-SpO₂ scenarios',10.2,C['text2'])
shape(s,10.56,4.65,2.09,1.34,C['amberbg'],None)
text(s,10.76,4.84,1.73,0.23,'CADENCE NOTE',8,C['amber'],True)
text(s,10.76,5.18,1.73,0.63,'Firmware constant is 60 s today; source comment says 2 s. Reconcile before pilot.',8.5,C['text2'])
text(s,0.65,6.48,11.6,0.27,'MQTT uses TLS configuration in the bridge; production firmware certificate verification and secrets handling need hardening.',9,C['muted'])

# 10 — Product map
s = new_slide('09  /  PLATFORM', 'A caregiver workflow—not just a chart—anchors the product.',
              'Implemented dashboard areas span observation, action, history and operations.', page=10)
modules=[
 ('LIVE OVERVIEW','Vitals + sparklines • IMU/activity • steps • connectivity status • event feed',C['purple']),
 ('ALERTS + INSIGHTS','Severity, trigger readings, explanations, recommendations, acknowledge / resolve',C['red']),
 ('SAFETY + SOS','Incident timeline, inactivity monitor, GPS map, contact selection, dispatch log',C['amber']),
 ('TRENDS','HRV, temperature / SpO₂, stress / HR, 14-day activity, 24h–30d ranges',C['lavender']),
 ('MANAGEMENT','Contacts CRUD, device / MQTT pairing, thresholds, wearer profile, care team',C['green']),
 ('SETTINGS + UX','Caregiver / admin roles, persistent session, theme, simulator and system controls',C['blue'])]
for i,(t,b,ac) in enumerate(modules):
    x=0.62+(i%3)*4.10; y=2.03+(i//3)*1.92
    card(s,x,y,3.83,1.64,t,b,ac,body_size=10.2,title_size=11.5)
shape(s,0.62,6.00,12.05,0.57,C['surface3'],None)
text(s,0.85,6.16,11.52,0.20,'Experience details: responsive layout • dark / light mode • empty states • skeleton loading • confirmation dialogs • optimistic updates',9.2,C['purple'],True,PP_ALIGN.CENTER)

# 11 — Explainable AI
s = new_slide('10  /  INTELLIGENCE', 'Personal calibration makes the live engine more transparent than a black box.',
              'Current detection uses equations + two-tier rules. Legacy ML assets remain in the repository, but are not the live engine.', page=11, title_size=23)
shape(s,0.62,2.02,7.40,4.37,C['surface'],C['border'])
text(s,0.91,2.29,5.8,0.25,'PHYSIOLOGY-INSPIRED SIGNAL FEATURES',9,C['purple'],True)
features=[
 ('Robust personal deviation','Median + MAD z-score; less sensitive to outliers.'),
 ('EDA tonic / phasic split','Slow baseline separates skin conductance level from response.'),
 ('Composite stress evidence','GSR change + HRV drop + HR rise at rest; terms stay inspectable.'),
 ('Thermal trend','Skin-to-core estimate calibrated against a guided reference; slope matters.'),
 ('Hypoxic burden','Sustained desaturation accumulates; one noisy sample is not the whole story.'),
 ('Fall motion sequence','Free-fall-like change → impact → rotation → stillness; tune to device data.')]
for i,(a,b) in enumerate(features):
    yy=2.76+i*0.54
    shape(s,0.95,yy+0.03,0.12,0.12,C['lavender'],None,True,kind=MSO_SHAPE.OVAL)
    text(s,1.20,yy,2.15,0.22,a,10,C['ink'],True)
    text(s,3.40,yy,4.18,0.37,b,9.0,C['text2'])
shape(s,8.31,2.02,4.36,4.37,C['dark'],None)
text(s,8.63,2.34,3.68,0.23,'CALIBRATION LOOP',9,C['lavender2'],True)
for i,(h,b) in enumerate([('Auto baseline','Quiet-gated EWMA from incoming readings'),('Auto-fit','Robust percentiles from recent history'),('Guided references','Thermometer, pulse-oximeter, resting HR / HRV'),('Manual override','Editable threshold and baseline controls')]):
    yy=2.82+i*0.72
    text(s,8.63,yy,3.60,0.22,h,11,C['white'],True)
    text(s,8.63,yy+0.27,3.57,0.35,b,8.8,'C5BFE8')
pill(s,8.64,5.86,2.06,'WARMING → ACTIVE',C['darkcard'],C['lavender2'],0.28,7.6,C['darkborder'])
text(s,0.65,6.60,11.8,0.20,'Reference sources are tracked (prior / auto / guided / manual); confidence increases with data volume. Baseline snapshots over time are future work.',8.5,C['muted'])

# 12 — Alert coverage
s = new_slide('11  /  DETECTION', 'A broad event set supports useful triage—not diagnosis.',
              'Alerts include severity, evidence readings, plain-language explanation and suggested next steps.', page=12)
conditions=[
 ('Fall detected','IMU impact / rotation / stillness sequence','Critical'),
 ('Low oxygen','SpO₂ floor + cumulative desaturation burden','High / urgent'),
 ('Fever','Temperature level + estimated core / rise over time','Medium / high'),
 ('High stress','EDA response + HRV drop + HR context','Medium'),
 ('Panic / acute distress','Fast HR + stress spike + motion context','High'),
 ('Tachycardia / bradycardia','HR outside wearer-specific range','Configurable'),
 ('Fatigue','Sustained low HRV / recovery pattern','Medium'),
 ('Inactivity','No meaningful movement beyond configured interval','Medium'),
 ('Emergency SOS','Caregiver / wearer-initiated action','Critical')]
# table
shape(s,0.62,2.00,12.05,0.44,C['purple'],None,True)
text(s,0.86,2.11,2.25,0.18,'EVENT',8,C['white'],True)
text(s,3.13,2.11,6.05,0.18,'SIGNAL / PATTERN',8,C['white'],True)
text(s,10.06,2.11,2.12,0.18,'SEVERITY',8,C['white'],True)
for i,(name,logic,sev) in enumerate(conditions):
    yy=2.55+i*0.42
    if i%2==0: shape(s,0.62,yy-0.03,12.05,0.40,C['surface'],None,True)
    text(s,0.86,yy,2.20,0.20,name,9.2,C['ink'],True)
    text(s,3.13,yy,6.55,0.22,logic,9.0,C['text2'])
    color=C['red'] if ('Critical' in sev or 'urgent' in sev) else (C['amber'] if ('High' in sev or 'Medium' in sev) else C['purple'])
    text(s,10.06,yy,2.1,0.20,sev,8.8,color,True)
text(s,0.66,6.49,11.7,0.22,'A rules-based engine can flag a pattern; sensor confirmation, symptom check and care-team judgement remain essential.',8.7,C['muted'],False,italic=True)

# 13 — Two-tier safety
s = new_slide('12  /  SAFETY MODEL', 'Personalization may add sensitivity; it must never erase a safety floor.',
              'Each condition can trigger when either the fixed clinical tier or the wearer-specific tier fires.', page=13, title_size=24)
# left/right cards
shape(s,0.62,2.08,5.42,2.72,C['surface'],C['border'])
pill(s,0.90,2.34,1.26,'TIER 1',C['redbg'],C['red'],0.28,8)
text(s,0.90,2.82,4.75,0.34,'Fixed safety boundaries',17,C['ink'],True)
bullet_list(s,0.92,3.39,4.78,[
 'Non-editable floors in the live engine',
 'Illustrative examples: low SpO₂, fever-range estimate, extreme resting HR, fall impact',
 'Designed to prevent unsafe baseline drift'],10.3,gap=0.41,dot_color=C['red'])
shape(s,7.30,2.08,5.37,2.72,C['surface'],C['border'])
pill(s,7.58,2.34,1.26,'TIER 2',C['lavender2'],C['purple'],0.28,8)
text(s,7.58,2.82,4.72,0.34,'Personal deviation rules',17,C['ink'],True)
bullet_list(s,7.60,3.39,4.77,[
 'Per-wearer median / MAD or σ-based thresholds',
 'Trend-aware deviations: HRV drop, stress, sustained oxygen burden',
 'Can be calibrated and tuned as evidence grows'],10.2,gap=0.41,dot_color=C['lavender'])
# OR gate
line(s,6.04,3.42,6.52,3.42,C['red'],2)
line(s,7.30,3.42,6.82,3.42,C['lavender'],2)
shape(s,6.30,3.11,0.52,0.52,C['purple'],None,True,kind=MSO_SHAPE.OVAL)
text(s,6.30,3.26,0.52,0.18,'OR',9,C['white'],True,PP_ALIGN.CENTER)
shape(s,1.45,5.33,10.45,0.79,C['dark'],None)
text(s,1.76,5.56,9.82,0.24,'TIER 1 fires  OR  TIER 2 fires   →   create an evidence-linked alert for a human to review',12,C['white'],True,PP_ALIGN.CENTER)
text(s,0.68,6.52,11.8,0.22,'Safety boundaries shown in software are engineering safeguards, not clinically validated thresholds for any individual.',8.8,C['muted'],False,italic=True)

# 14 — Alert workflow
s = new_slide('13  /  RESPONSE', 'The alert lifecycle connects detection to accountable action.',
              'Most of the caregiver workflow is implemented; external notification delivery is still a production integration.', page=14)
stages=[('DETECT','Candidate pattern\n+ sensor values',C['purple']),('PRIORITIZE','Severity\n+ why it fired',C['red']),('REVIEW','Acknowledge\n+ confirm context',C['amber']),('ESCALATE','Select trusted\ncontact / channel',C['lavender']),('CLOSE LOOP','Log delivery,\nresolve incident',C['green'])]
for i,(a,b,ac) in enumerate(stages):
    x=0.62+i*2.55
    shape(s,x,2.55,2.14,1.68,C['surface'],C['border'])
    shape(s,x+0.20,2.76,0.43,0.43,ac,None,True,kind=MSO_SHAPE.OVAL)
    text(s,x+0.20,2.87,0.43,0.18,str(i+1),9,C['white'],True,PP_ALIGN.CENTER)
    text(s,x+0.20,3.34,1.73,0.24,a,9,ac,True)
    text(s,x+0.20,3.68,1.75,0.44,b,9.5,C['text2'])
    if i<4: line(s,x+2.15,3.39,x+2.47,3.39,C['lavender'],1.6)
shape(s,0.62,4.70,5.82,1.37,C['greenbg'],None)
text(s,0.90,4.93,5.23,0.25,'IN PRODUCT TODAY',9,C['green'],True)
text(s,0.90,5.31,5.18,0.48,'Confirm-to-SOS • acknowledge / resolve • contact CRUD • selectable SMS / call / app labels • dispatch history',10.0,C['text2'])
shape(s,6.72,4.70,5.95,1.37,C['amberbg'],None)
text(s,7.00,4.93,5.31,0.25,'INTEGRATION BEFORE REAL-WORLD USE',9,C['amber'],True)
text(s,7.00,5.31,5.21,0.48,'Verified SMS / voice / push provider • delivery receipts • retries • offline fallback • consent + escalation policy',10.0,C['text2'])
text(s,0.65,6.51,11.9,0.20,'Current dispatch endpoint writes a workflow record in the prototype; it does not place an actual call or send production SMS.',9,C['muted'],True)

# 15 — Safety workspace mockup
s = new_slide('14  /  SAFETY EXPERIENCE', 'The safety view keeps incident, location and next action together.',
              'A calmer response surface reduces the need to jump between tools during an event.', page=15)
# UI frame
shape(s,0.62,1.99,12.05,4.50,C['surface'],C['border'])
shape(s,0.62,1.99,12.05,0.42,C['purple'],None,True)
text(s,0.87,2.10,3.40,0.18,'SAFETY  /  MARGARET THOMPSON',8,C['white'],True)
pill(s,10.82,2.08,1.46,'STATUS: REVIEW',C['amberbg'],C['amber'],0.23,7)
# timeline left
text(s,0.91,2.72,3.2,0.24,'INCIDENT TIMELINE',8,C['muted'],True)
line(s,1.02,3.17,1.02,5.77,C['border'],2)
entries=[('18:42','Fall-like motion signature','Critical / active',C['red']),('18:44','Caregiver acknowledged','Follow-up requested',C['amber']),('18:48','Contact selected','Dispatch record saved',C['green'])]
for i,(t,a,b,ac) in enumerate(entries):
    yy=3.15+i*0.87
    shape(s,0.91,yy,0.22,0.22,ac,None,True,kind=MSO_SHAPE.OVAL)
    text(s,1.34,yy-0.01,0.62,0.18,t,8,ac,True)
    text(s,2.00,yy-0.02,2.10,0.22,a,9.0,C['ink'],True)
    text(s,2.00,yy+0.25,2.15,0.30,b,8,C['muted'])
# map mock right
shape(s,4.72,2.73,5.05,3.20,'F5F4FA',C['border'])
# abstract roads
for x in [5.4,6.35,7.3,8.25,9.15]: line(s,x,2.80,x+0.25,5.88,'E2DFEF',2)
for y in [3.30,4.00,4.70,5.40]: line(s,4.78,y,9.68,y+0.18,'E2DFEF',2)
line(s,5.18,5.58,6.10,5.00,C['lavender'],2.3)
line(s,6.10,5.00,7.14,4.45,C['lavender'],2.3)
line(s,7.14,4.45,8.05,4.56,C['lavender'],2.3)
shape(s,7.84,4.25,0.33,0.33,C['red'],None,True,kind=MSO_SHAPE.OVAL)
text(s,7.91,4.32,0.20,0.16,'+',10,C['white'],True,PP_ALIGN.CENTER)
text(s,4.96,5.67,4.58,0.15,'ILLUSTRATIVE MAP / LOCATION MARKER',7.2,C['muted'],True,PP_ALIGN.CENTER)
# action panel
shape(s,10.00,2.73,2.37,3.20,C['surface2'],C['border'])
text(s,10.22,2.97,1.93,0.22,'NEXT ACTION',8,C['muted'],True)
text(s,10.22,3.35,1.94,0.60,'Check on\nwearer',15,C['ink'],True)
text(s,10.22,4.12,1.93,0.68,'Select trusted contacts and an approved channel.',8.8,C['text2'])
pill(s,10.22,5.10,1.92,'OPEN DISPATCH',C['purple'],C['white'],0.31,7.6)
text(s,0.91,6.10,11.4,0.18,'GPS map / event markers and dispatch history exist in the prototype; map accuracy, provider delivery and indoor location require validation.',8.4,C['muted'])

# 16 — Trends
s = new_slide('15  /  TRENDS', 'Longitudinal context turns one alert into a useful history.',
              'Interactive charts, range filters, clinical overlays and daily activity summaries are implemented.', page=16)
# metric cards
for i,(lab,val,sub) in enumerate([('HRV','42 ms','24h avg'),('TEMP','36.6 °C','24h avg'),('SpO₂','96.8 %','24h avg'),('STRESS','0.29','24h avg')]):
    x=0.62+i*3.05
    shape(s,x,2.00,2.78,0.91,C['surface'],C['border'])
    text(s,x+0.18,2.16,2.33,0.15,lab,7.5,C['muted'],True)
    text(s,x+0.18,2.39,1.38,0.33,val,18,C['ink'],True)
    text(s,x+1.55,2.52,1.02,0.17,sub,7.5,C['muted'])
# chart panels with editable line plots
charts=[('HRV HISTORY',[(0,0.56),(0.14,0.43),(0.28,0.48),(0.42,0.33),(0.57,0.40),(0.72,0.24),(0.88,0.31)]),
        ('TEMP + SpO₂',[(0,0.55),(0.16,0.50),(0.32,0.54),(0.48,0.38),(0.64,0.43),(0.82,0.33),(1,0.35)]),
        ('STRESS + HEART RATE',[(0,0.62),(0.16,0.56),(0.30,0.40),(0.46,0.46),(0.60,0.26),(0.77,0.38),(1,0.22)])]
for i,(title_,values) in enumerate(charts):
    x=0.62+i*4.10
    shape(s,x,3.20,3.82,2.33,C['surface'],C['border'])
    text(s,x+0.20,3.38,3.42,0.18,title_,8,C['muted'],True)
    for j in range(3): line(s,x+0.24,3.82+j*0.44,x+3.55,3.82+j*0.44,'ECEAF4',0.6)
    pts=[(x+0.28+px*3.2,5.20-py*1.2) for px,py in values]
    for a,b in zip(pts,pts[1:]): line(s,a[0],a[1],b[0],b[1],C['lavender'],1.8)
    line(s,x+0.26,4.31,x+3.55,4.31,C['amber'],0.8,MSO_LINE_DASH_STYLE.DASH)
# ranges + activity
pill(s,0.67,5.83,0.73,'24 h',C['purple'],C['white'],0.27,7.6)
for j,t in enumerate(['48 h','7 d','30 d']): pill(s,1.49+j*0.74,5.83,0.70,t,C['surface3'],C['muted'],0.27,7.4)
text(s,4.75,5.88,7.85,0.18,'Also: 14-day step / activity bars • today’s motion mix • average and delta tiles • time-series tooltips',8.8,C['text2'])
text(s,0.66,6.48,11.8,0.18,'Charts shown are vector mockups with illustrative data; source dashboard renders dynamic readings and selected time windows.',8.5,C['muted'],False,italic=True)

# 17 — Management and roles
s = new_slide('16  /  OPERATIONS', 'Caregiver and device administration is part of the core product.',
              'Operational controls are built into the current management views rather than left to engineering.', page=17)
ops=[
 ('CONTACTS','Create / edit / remove trusted contacts\nPriority order + dispatch eligibility\nRelationship and communication details',C['purple']),
 ('DEVICE + MQTT','Pair / unpair device and serial\nBroker, topic, TLS configuration\nConnection test with stage-level feedback',C['lavender']),
 ('PERSONALIZATION','Per-patient alert thresholds\nCalibration and reference history\nWearer profile / conditions / notes',C['green']),
 ('CARE TEAM','Admin user create / edit / remove\nAdmin and caregiver role gates\nCaregiver read-only team access',C['amber'])]
for i,(t,b,ac) in enumerate(ops):
    x=0.62+(i%2)*6.10; y=2.02+(i//2)*1.97
    card(s,x,y,5.82,1.67,t,b,ac,body_size=10.4,title_size=12)
shape(s,0.62,6.16,12.05,0.47,C['surface3'],None)
text(s,0.84,6.29,11.55,0.18,'Caregiver UX: persistent login • quick SOS confirmation • dark / light theme • accessible empty states • admin-only actions gated',8.8,C['purple'],True,PP_ALIGN.CENTER)

# 18 — Privacy/security
s = new_slide('17  /  TRUST', 'Trust is a product requirement—and a current engineering workstream.',
              'Prototype controls provide a foundation; production deployment needs a formal privacy and security program.', page=18, title_size=24)
shape(s,0.62,2.02,5.83,3.93,C['surface'],C['border'])
text(s,0.91,2.29,5.1,0.24,'IN PLACE IN THE SOFTWARE PROTOTYPE',9,C['green'],True)
bullet_list(s,0.92,2.81,5.05,[
 'PBKDF2 password hashing; HMAC-signed sessions',
 'Admin / caregiver role-based permissions',
 'Device-key protected telemetry ingestion',
 'MQTT bridge supports TLS certificate validation',
 'Audit fields for alert acknowledgement / resolution'],10.1,gap=0.56,dot_color=C['green'])
shape(s,6.75,2.02,5.92,3.93,C['amberbg'],None)
text(s,7.04,2.29,5.25,0.24,'BEFORE PILOT OR COMMERCIAL USE',9,C['amber'],True)
bullet_list(s,7.05,2.81,5.12,[
 'Remove source-embedded device credentials; rotate keys',
 'Replace firmware TLS bypass with certificate validation',
 'Encrypt data at rest / in transit; multi-tenant isolation',
 'Consent, purpose limitation, retention and deletion policy',
 'Threat model, independent penetration test, incident plan'],9.9,gap=0.56,dot_color=C['amber'])
text(s,0.66,6.29,11.8,0.48,'Regulatory scope depends on intended use and claims. Do not market diagnosis or treatment until counsel, quality systems and the relevant regulator define the pathway.',9.2,C['purple'],True)

# 19 — Differentiation
s = new_slide('18  /  POSITIONING', 'The wedge is explainability plus action—not “AI” as a label.',
              'A credible advantage must be earned through calibration quality, workflow fit and trustworthy field evidence.', page=19)
# comparison grid
shape(s,0.62,2.02,12.05,3.56,C['surface'],C['border'])
# columns headers
text(s,0.88,2.27,2.35,0.20,'CAPABILITY',8,C['muted'],True)
text(s,3.50,2.27,3.35,0.20,'GENERIC TRACKER PATTERN',8,C['muted'],True)
text(s,7.24,2.27,4.92,0.20,'NEUROLINK DESIGN INTENT',8,C['purple'],True)
rows=[
 ('Context','Single-metric reading','Multiple signals + motion / personal history'),
 ('Adaptation','Fixed one-size threshold','Quiet-gated personal baseline + guided reference'),
 ('Explainability','Score with limited rationale','Expose sensor evidence and rule contribution'),
 ('Safety logic','One threshold decides','Fixed floors OR wearer-specific deviations'),
 ('Response','Wearer-only notification','Caregiver workflow, GPS context, event history')]
for i,(a,b,c_) in enumerate(rows):
    yy=2.73+i*0.53
    line(s,0.84,yy-0.08,12.38,yy-0.08,C['border'],0.6)
    text(s,0.90,yy,2.30,0.22,a,9.6,C['ink'],True)
    text(s,3.50,yy,3.25,0.25,b,9.2,C['muted'])
    text(s,7.24,yy,4.84,0.25,c_,9.2,C['text2'])
shape(s,0.62,5.86,12.05,0.57,C['dark'],None)
text(s,0.89,6.03,11.48,0.21,'Defensibility hypothesis: high-quality longitudinal data + low-noise alerting + trusted care workflows + distribution partnerships.',9.7,C['white'],True,PP_ALIGN.CENTER)
text(s,0.66,6.60,11.7,0.17,'These are product design choices, not a proven competitive superiority claim. Competitor and IP diligence remain open.',8.4,C['muted'],False,italic=True)

# 20 — Validation/readiness
s = new_slide('19  /  READINESS', 'Software quality is evidenced. Clinical and hardware quality are not yet established.',
              'The immediate goal is an instrumented, safe, measurable pilot—not a broad launch.', page=20, title_size=23)
# left proof column
shape(s,0.62,2.02,5.45,3.98,C['surface'],C['border'])
text(s,0.91,2.30,4.80,0.23,'REPOSITORY EVIDENCE',9,C['purple'],True)
for i,(v,l) in enumerate([('141/141','end-to-end / backend / UI QA'),('74/74 × 2','equations + calibration suite'),('17/17','calibration DOM checks'),('Seeded demo','7-day telemetry, alerts, GPS / history')]):
    yy=2.78+i*0.69
    text(s,0.93,yy,1.46,0.28,v,13,C['ink'],True)
    text(s,2.47,yy+0.02,3.22,0.34,l,9.3,C['text2'])
    if i<3: line(s,0.93,yy+0.47,5.72,yy+0.47,C['border'],0.6)
text(s,0.93,5.63,4.80,0.20,'Strong software test coverage ≠ real-world efficacy.',8.5,C['amber'],True)
# right gaps
shape(s,6.42,2.02,6.25,3.98,C['amberbg'],None)
text(s,6.72,2.30,5.60,0.23,'GATES TO PILOT',9,C['amber'],True)
bullet_list(s,6.73,2.78,5.45,[
 'Bench reference tests: HR / SpO₂ / temperature / HRV',
 'Signal quality + motion / fit / skin-tone subgroup analysis',
 'Battery life, charge habits, drop / water / wear testing',
 'False alerts per wearer-day, sensitivity, latency, adherence',
 'Real broker validation; notification provider receipts',
 'Privacy impact, consent, safety-case and clinical review'],9.4,gap=0.49,dot_color=C['amber'])
text(s,0.67,6.37,11.7,0.28,'No prospective clinical validation, regulatory clearance, production manufacturing qualification, customer revenue or pilot outcomes are claimed in this deck.',8.8,C['muted'],True)

# 21 — Business model
s = new_slide('20  /  BUSINESS MODEL', 'A blended device + recurring software model aligns value with ongoing care.',
              'Use pilots to test pricing, attach rate, retention and channel economics.', page=21)
models=[
 ('DEVICE SALE','Illustrative $149','Band + charging kit; upfront revenue; margin depends on verified BOM, yield and warranty.',C['purple']),
 ('FAMILY PLAN','Illustrative $12 / mo','Caregiver dashboard, alerts, longer history and trusted-contact workflow; household subscription.',C['lavender']),
 ('CARE PARTNER','Illustrative $7–10 / wearer / mo','Volume license, team workflow and reporting; onboarding / integration can be paid separately.',C['green'])]
for i,(a,b,c_,ac) in enumerate(models):
    x=0.62+i*4.10
    shape(s,x,2.12,3.83,2.64,C['surface'],C['border'])
    text(s,x+0.23,2.40,3.31,0.24,a,9,ac,True)
    text(s,x+0.23,2.88,3.31,0.51,b,20,C['ink'],True)
    text(s,x+0.23,3.57,3.31,0.91,c_,10.4,C['text2'])
shape(s,0.62,5.08,12.05,1.17,C['dark'],None)
text(s,0.90,5.31,2.20,0.22,'EARLY REVENUE BRIDGE',8.5,C['lavender2'],True)
text(s,3.10,5.23,9.15,0.46,'Paid pilots / setup fees  →  device + family subscriptions  →  care-organization contracts  →  partner distribution',11,C['white'],True,PP_ALIGN.CENTER,MSO_ANCHOR.MIDDLE)
text(s,0.66,6.50,11.9,0.21,'Pricing is an assumption to validate with buyers, not a published offer. Reimbursement is excluded from the base case.',8.7,C['muted'],False,italic=True)

# 22 — GTM
s = new_slide('21  /  GO TO MARKET', 'Prove the workflow in a narrow pilot before buying broad consumer reach.',
              'Success means a trusted response loop, repeat use and a viable cost to serve.', page=22)
phases=[
 ('PILOT DESIGN','0–3 months','2–3 partners\n100–150 target wearers\nCaregiver + wearer interviews',C['purple']),
 ('CONTROLLED PILOT','3–9 months','Safety-reviewed protocol\nMeasure adherence + alerts\nWeekly service feedback',C['lavender']),
 ('REPEATABLE SALES','9–18 months','Convert paid sites\nFamily plan tests\nChannel / support playbook',C['green'])]
for i,(a,b,c_,ac) in enumerate(phases):
    x=0.62+i*4.10
    shape(s,x,2.18,3.83,2.33,C['surface'],C['border'])
    pill(s,x+0.22,2.42,1.14,b,C['surface3'],ac,0.27,7.5)
    text(s,x+0.22,2.93,3.38,0.28,a,11.5,C['ink'],True)
    text(s,x+0.22,3.39,3.35,0.86,c_,10.4,C['text2'])
    if i<2: line(s,x+3.84,3.32,x+4.04,3.32,C['lavender'],1.5)
shape(s,0.62,4.92,12.05,1.29,C['surface2'],C['border'])
text(s,0.90,5.17,2.30,0.23,'PILOT SCORECARD',8.5,C['purple'],True)
text(s,3.10,5.10,9.12,0.60,'Wear-days / active days  •  battery days  •  user retention  •  alert PPV / false alerts per day  •  time to acknowledgement  •  contact delivery success  •  willingness to pay',10.0,C['text2'],False,PP_ALIGN.CENTER,MSO_ANCHOR.MIDDLE)
text(s,0.67,6.49,11.7,0.18,'Target channel: family caregiving groups, home-care operators and senior-living pilot partners; no signed partner is represented here.',8.4,C['muted'],False,italic=True)

# 23 — Financial plan
s = new_slide('22  /  FINANCIAL PLAN', 'Illustrative base case: build recurring accounts as device placements scale.',
              'Management planning scenario only; assumptions are not actual results, commitments or investment guidance.', page=23, title_size=23)
# chart background
shape(s,0.62,2.03,7.30,3.95,C['surface'],C['border'])
text(s,0.91,2.27,6.6,0.21,'REVENUE BUILD  •  $M',8.4,C['muted'],True)
vals=[0.1855,0.863,2.926,7.373,15.134]
labels=['Y1','Y2','Y3','Y4','Y5']
maxv=16
for i,(v,lab) in enumerate(zip(vals,labels)):
    x=1.18+i*1.25; h=2.60*(v/maxv)
    shape(s,x,5.43-h,0.64,h,C['purple'] if i<3 else C['lavender'],None,True)
    text(s,x-0.18,5.43-h-0.28,1.00,0.20,f'${v:.2f}M',8.2,C['ink'],True,PP_ALIGN.CENTER)
    text(s,x-0.02,5.53,0.68,0.18,lab,8,C['muted'],True,PP_ALIGN.CENTER)
line(s,1.00,5.44,7.63,5.44,C['border'],0.8)
# right table
shape(s,8.18,2.03,4.49,3.95,C['dark'],None)
text(s,8.46,2.29,3.92,0.20,'PLANNING INPUTS',8.5,C['lavender2'],True)
inputs=[('Devices sold','500 → 50,000'),('Avg paid subs','250 → 36,000'),('Device ASP','$149'),('Subscription','$12 / month'),('B2B revenue','$75k → $2.5M')]
for i,(a,b) in enumerate(inputs):
    yy=2.78+i*0.49
    text(s,8.48,yy,1.68,0.21,a,8.7,'C5BFE8')
    text(s,10.16,yy,2.17,0.21,b,9,C['white'],True,PP_ALIGN.RIGHT)
    if i<4: line(s,8.47,yy+0.31,12.31,yy+0.31,C['darkborder'],0.7)
text(s,8.47,5.36,3.85,0.35,'No reimbursement; excludes taxes.',8.4,C['gray'],False,italic=True)
# revenue table note
text(s,0.67,6.16,11.9,0.23,'Revenue mix: device ASP × units + average paid subscriptions × $144 annual ARPU + B2B pilot / license revenue.',9.2,C['text2'],True)
text(s,0.67,6.54,11.8,0.19,'Illustrative 5-year scenario (USD). Actual prices, churn, hardware yield, CAC, returns and partner sales cycles remain unvalidated.',8.5,C['muted'],False,italic=True)

# 24 — Unit economics
s = new_slide('23  /  ECONOMICS', 'The model can reach operating break-even—but only if hardware and retention assumptions hold.',
              'Base-case contribution improves with procurement scale and recurring subscriptions.', page=24, title_size=22)
# table
shape(s,0.62,2.01,12.05,3.34,C['surface'],C['border'])
headers=['$M except margin','Y1','Y2','Y3','Y4','Y5']
xs=[0.88,5.16,6.55,7.93,9.31,10.69]
for x,hdr in zip(xs,headers): text(s,x,2.28,1.22,0.20,hdr,8.3,C['muted'] if x!=0.88 else C['purple'],True,PP_ALIGN.RIGHT if x>1 else PP_ALIGN.LEFT)
rows=[('Revenue','0.19','0.86','2.93','7.37','15.13'),
      ('Gross profit','0.10','0.49','1.84','4.83','10.25'),
      ('Gross margin','52%','57%','63%','66%','68%'),
      ('Operating expense','0.70','1.25','2.20','3.60','5.80'),
      ('EBITDA (planning)','(0.60)','(0.76)','(0.36)','1.23','4.45')]
for i,row in enumerate(rows):
    yy=2.76+i*0.46
    if i==4: shape(s,0.82,yy-0.06,11.48,0.41,C['greenbg'],None)
    text(s,0.90,yy,3.85,0.20,row[0],9.2,C['ink'],i in (0,4))
    for j,v in enumerate(row[1:]): text(s,5.16+j*1.38,yy,1.15,0.20,v,9.0,C['green'] if i==4 and j>=3 else C['text2'],i in (0,4),PP_ALIGN.RIGHT)
    if i<4: line(s,0.89,yy+0.30,12.27,yy+0.30,C['border'],0.5)
shape(s,0.62,5.70,5.82,0.67,C['surface2'],C['border'])
text(s,0.85,5.83,1.57,0.20,'DEVICE CONTRIBUTION',7.5,C['muted'],True)
text(s,2.54,5.81,3.50,0.24,'$149 ASP − $90 → $52 unit COGS',10.0,C['purple'],True)
shape(s,6.73,5.70,5.94,0.67,C['surface2'],C['border'])
text(s,6.97,5.83,1.66,0.20,'SUBSCRIPTION',7.5,C['muted'],True)
text(s,8.66,5.81,3.70,0.24,'$12/mo; model assumes 80% gross margin',9.8,C['purple'],True)
text(s,0.66,6.58,11.8,0.17,'Before financing, tax, depreciation, working capital, CAC, fulfillment, warranty and returns. Break-even timing is highly sensitive to adoption and support costs.',8.2,C['muted'],False,italic=True)

# 25 — Roadmap
s = new_slide('24  /  PRODUCT ROADMAP', 'Build next for reliability, usability and partner readiness.',
              'Future features are hypotheses and roadmap candidates—not committed or shipped functionality.', page=25, title_size=24)
road=[
 ('0–6 MONTHS','Hardware + safety foundation','Fix firmware cadence / time source\nCertificate validation + secret rotation\nBench-test sensors + signal quality\nBattery / enclosure / wearability iteration',C['purple']),
 ('6–18 MONTHS','Pilot-ready product','Offline buffering + reconnect UX\nWearer companion app + haptics / SOS\nReal SMS / voice / push delivery receipts\nCalibration history / drift reports',C['lavender']),
 ('18–36 MONTHS','Scale + interoperability','Multi-tenant care-team portal\nRole / consent / audit / retention controls\nFHIR / provider integrations (partner-led)\nValidated ML / on-device inference if justified',C['green'])]
for i,(date,title_,body,ac) in enumerate(road):
    x=0.62+i*4.10
    shape(s,x,2.13,3.83,3.65,C['surface'],C['border'])
    pill(s,x+0.23,2.38,1.23,date,C['surface3'],ac,0.28,7.5)
    text(s,x+0.23,2.88,3.36,0.51,title_,14,C['ink'],True)
    line(s,x+0.23,3.54,x+3.52,3.54,C['border'],0.8)
    bullet_list(s,x+0.24,3.78,3.30,body.split('\n'),9.7,gap=0.45,dot_color=ac)
shape(s,0.62,6.03,12.05,0.53,C['purple'],None)
text(s,0.88,6.19,11.49,0.20,'Roadmap order is evidence-led: safety + reliability first, feature breadth second.',9.5,C['white'],True,PP_ALIGN.CENTER)

# 26 — Risks
s = new_slide('25  /  EXECUTION RISKS', 'The risks are manageable only if we name them early.',
              'A pilot should be designed to test failure modes, not merely demonstrate the happy path.', page=26)
riskrows=[
 ('Measurement bias / noise','Reference-instrument bench testing; per-signal quality checks; subgroup review',C['red']),
 ('False alerts / missed events','Pre-register endpoints; shadow mode; human review; alert audit and tuning',C['amber']),
 ('Battery / wear adherence','Daily wear-time telemetry; charger usability tests; field-replaceable design',C['purple']),
 ('Emergency delivery failure','Provider receipts, retries, escalation fallback and explicit non-emergency disclaimers',C['lavender']),
 ('Security / health-data privacy','Credential rotation, verified TLS, threat model, least privilege, retention controls',C['green']),
 ('Regulatory / liability ambiguity','Define intended use; avoid diagnostic claims; clinical and legal review before pilots',C['blue'])]
for i,(a,b,ac) in enumerate(riskrows):
    yy=2.02+i*0.65
    if i%2==0: shape(s,0.62,yy-0.03,12.05,0.57,C['surface'],None,True)
    shape(s,0.82,yy+0.08,0.12,0.12,ac,None,True,kind=MSO_SHAPE.OVAL)
    text(s,1.10,yy+0.01,3.00,0.23,a,9.5,C['ink'],True)
    text(s,4.14,yy+0.01,8.02,0.34,b,9.2,C['text2'])
text(s,0.66,6.30,11.8,0.25,'Product guardrail: wellness / safety support for human review; not a substitute for emergency services or professional medical advice.',9.2,C['purple'],True)

# 27 — Funding and pilot ask
s = new_slide('26  /  NEXT STEP', 'A focused pre-seed plan can fund the evidence—not just the feature list.',
              'Illustrative raise scenario: $1.5M for ~18 months to reach a validated, pilot-ready milestone.', page=27, title_size=24)
shape(s,0.62,2.03,4.27,3.93,C['dark'],None)
text(s,0.98,2.40,3.48,0.22,'ILLUSTRATIVE FUNDING TARGET',8.5,C['lavender2'],True)
text(s,0.98,2.91,3.46,0.78,'$1.5M',38,C['white'],True)
text(s,0.98,3.81,3.47,0.29,'~18 months runway plan',12,'C5BFE8',True)
line(s,0.98,4.34,4.49,4.34,C['darkborder'],1)
text(s,0.98,4.58,3.40,0.97,'Outcome: safer hardware cadence, credible sensor evidence, real delivery integrations and partner pilots with measurable endpoints.',10,C['white'])
# Use of funds bars
shape(s,5.31,2.03,7.36,3.93,C['surface'],C['border'])
text(s,5.61,2.36,6.57,0.24,'USE OF FUNDS  •  PLANNING ALLOCATION',8.5,C['purple'],True)
alloc=[('Product + hardware verification',35,C['purple']),('Pilot / evidence / safety validation',25,C['lavender']),('Engineering + data platform',20,C['green']),('Security / privacy / regulatory',10,C['amber']),('Partner GTM + contingency',10,C['red'])]
for i,(a,pct,ac) in enumerate(alloc):
    yy=2.86+i*0.54
    text(s,5.62,yy,3.35,0.21,a,9.1,C['text2'],True)
    shape(s,8.97,yy+0.04,2.50,0.15,C['surface3'],None,True)
    shape(s,8.97,yy+0.04,2.50*pct/35,0.15,ac,None,True)
    text(s,11.58,yy-0.01,0.65,0.22,f'{pct}%',9.1,ac,True,PP_ALIGN.RIGHT)
text(s,5.62,5.70,6.47,0.18,'Milestone: 2–3 pilot partners • 100–150 target wearers • pre-agreed safety / usability metrics',8.4,C['muted'],True)
text(s,0.65,6.36,11.7,0.26,'Seeking pilot partners, aging-at-home operators, clinical validation advisors and aligned early-stage investors. Raise size / timing requires founder confirmation.',8.9,C['muted'],False,italic=True)

# 28 — Close
s = prs.slides.add_slide(blank); set_bg(s, C['dark'])
shape(s,0.68,0.55,0.17,0.17,C['lavender2'],None,True,kind=MSO_SHAPE.OVAL)
text(s,0.99,0.49,2.8,0.28,'NEUROLINK  /  WEAR',10,C['white'],True)
# abstract data signal
for i,hh in enumerate([0.55,0.90,0.38,1.22,0.73,1.54,0.48,1.06,0.67]):
    shape(s,8.18+i*0.43,2.11+(1.54-hh),0.12,hh,C['purple'] if i%2==0 else C['lavender'],None,True)
line(s,8.02,3.65,12.20,3.65,C['darkborder'],1)
text(s,0.72,1.77,7.12,1.36,'Make the signal\nuseful to someone.',34,C['white'],True,spacing=0.95)
text(s,0.77,3.50,6.75,0.65,'Wearable context. Explainable alerts. A safer path to human response.',15,'C5BFE8')
pill(s,0.78,4.76,2.02,'PILOT + VALIDATION',C['darkcard'],C['lavender2'],0.33,8.2,C['darkborder'])
pill(s,2.96,4.76,2.15,'PARTNER DISCUSSION',C['darkcard'],C['lavender2'],0.33,8.2,C['darkborder'])
text(s,0.78,6.48,6.2,0.26,'NeuroLink Wear  •  Investor / technology discussion  •  2026',9,C['gray'])
text(s,9.04,5.00,3.15,0.62,'Let’s prove it\nwith the people who need it.',15,C['white'],True,PP_ALIGN.RIGHT)

# Set document metadata.
prs.core_properties.title = 'NeuroLink Wear — Investor & Technology Pitch'
prs.core_properties.subject = 'Editable investor presentation covering product, technology, audience, business plan, financial scenario and roadmap'
prs.core_properties.author = 'NeuroLink Wear'
prs.core_properties.keywords = 'NeuroLink Wear, smart band, caregiver safety, wearable, investor deck'
prs.core_properties.comments = 'All visual elements are native editable PowerPoint text and shapes. Financials/pricing are explicitly illustrative assumptions.'

OUT.parent.mkdir(parents=True, exist_ok=True)
prs.save(OUT)
print(f'Wrote {OUT} ({len(prs.slides)} slides)')
