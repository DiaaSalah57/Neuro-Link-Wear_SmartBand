"""Deck section B — solution, product tour, hardware, architecture."""
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.dml import MSO_LINE_DASH_STYLE

from kit import *


def build(b):
    solution(b)
    product(b)
    hardware(b)
    architecture(b)


# ─────────────────────────────── 07 solution ───────────────────────────
def solution(b):
    s = slide(b.prs)
    n = b.num()
    header(s, "02 / Solution", "Sense, understand, respond — then learn the wearer",
           "A device, an intelligence layer and a care workflow that share one data spine.", num=n)

    steps = [
        ("1", "SENSE", "PPG, GSR, motion, temperature and GPS",
         "The band samples heart rate, blood oxygen, HRV, skin conductance, skin temperature and "
         "6-axis motion — plus location context — and publishes every 1–5 s over MQTT.", ACCENT),
        ("2", "UNDERSTAND", "Equations + calibration + language",
         "Raw signals become features, features become equations, equations become a condition — "
         "with a plain-language explanation of the evidence behind it.", ACCENT3),
        ("3", "RESPOND", "Explain, recommend, escalate",
         "The caregiver sees what changed, why it matters and what to do — and can reach a human "
         "with one action, with the whole event logged.", PURPLE),
    ]
    cw = gw(4, 0.24)
    evidence = ["25 Hz motion sampling · 1 Hz vitals · published every 1–5 s",
                "Evaluated on every reading at a ~2-second cadence",
                "notify → confirm → escalate → log, with GPS attached"]
    for i, (idx, tag, sub, body, col) in enumerate(steps):
        x = gx(i * 4)
        card(s, x, 1.78, cw, 2.96, fill=SURF)
        rect(s, x, 1.78, cw, 0.055, fill=col, radius=True, adj=0.5)
        oval(s, x + 0.26, 2.02, 0.42, 0.42, col)
        tx(s, x + 0.26, 2.09, 0.42, 0.30, idx, size=13, color=LIGHT, bold=True, align="c")
        tx(s, x + 0.80, 2.06, 2.4, 0.24, tag, size=11, color=col, bold=True, track=110)
        tx(s, x + 0.80, 2.30, 2.5, 0.24, sub, size=8, color=MUTED, line=1.15)
        hline(s, x + 0.26, 2.68, cw - 0.52, color=BORDER, lw=0.75)
        tx(s, x + 0.26, 2.84, cw - 0.52, 1.4, body, size=9.5, color=TEXT2, line=1.32)
        hline(s, x + 0.26, 4.24, cw - 0.52, color=BORDER, lw=0.5)
        tx(s, x + 0.26, 4.34, cw - 0.52, 0.3, evidence[i], size=8, color=col, line=1.2)
        if i < 2:
            ar = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x + cw + 0.03),
                                    Inches(3.14), Inches(0.18), Inches(0.16))
            ar.shadow.inherit = False
            ar.fill.solid(); ar.fill.fore_color.rgb = rgb(BORDER_STRONG)
            ar.line.fill.background()

    # closing loop band
    card(s, ML, 4.92, 6.62, 1.98, fill=SURF2)
    tx(s, ML + 0.26, 5.10, 6.0, 0.24, "THE LOOP CLOSES ON THE WEARER", size=8, color=MUTED, bold=True, track=120)
    tx(s, ML + 0.26, 5.38, 6.1, 0.62,
       "Every reading and every acknowledged alert feeds back into the wearer's personal baselines, "
       "so the next threshold is better than the last.", size=10, color=TEXT2, line=1.32)
    loop = [("Reading", ACCENT), ("Explained alert", ACCENT3), ("Human response", PURPLE),
            ("Baseline update", OK), ("Sharper alert", ACCENT2)]
    lx = ML + 0.26
    for i, (label, col) in enumerate(loop):
        w = 1.06
        chip(s, lx, 6.16, w, 0.34, label, fg=col, bgc=SURF, size=7.5, line=BORDER_STRONG, track=20)
        if i < len(loop) - 1:
            tx(s, lx + w, 6.16, 0.14, 0.34, "›", size=11, color=MUTED, align="c", anchor="m")
        lx += w + 0.14
    tx(s, ML + 0.26, 6.62, 6.1, 0.22,
       "The last step is the business model: it gets better the longer it is worn.",
       size=8, color=MUTED)

    card(s, 7.68, 4.92, 4.86, 1.98, fill=SURF)
    tx(s, 7.94, 5.10, 4.3, 0.24, "WHAT MAKES IT DIFFERENT", size=8, color=MUTED, bold=True, track=120)
    bullets(s, 7.94, 5.42, 4.3, 1.4,
            [("No dataset dependency. ", "Works on day one for a new wearer."),
             ("Same math on device and server. ", "The ESP32 can run the identical equations."),
             ("Explained, not scored. ", "Alerts name the signals and the deviation.")],
            size=9, dot="·", gap=5, line=1.26)
    footer(s, n, right="Solution")


# ─────────────────────────────── 08 product ────────────────────────────
def product(b):
    s = slide(b.prs)
    n = b.num()
    header(s, "02 / Product", "Six views, one caregiver workflow",
           "Shipped in the current build — a responsive, role-aware web app with live WebSocket data.", num=n)

    # ── dashboard wireframe
    dx, dy, dw, dh = ML, 1.74, 7.72, 4.2
    card(s, dx, dy, dw, dh, fill=SURF)
    rect(s, dx, dy, dw, 0.36, fill=SURF3, radius=True, adj=0.35)
    rect(s, dx, dy + 0.18, dw, 0.18, fill=SURF3, radius=False)
    dot(s, dx + 0.16, dy + 0.13, 0.10, BORDER_STRONG)
    dot(s, dx + 0.32, dy + 0.13, 0.10, BORDER_STRONG)
    dot(s, dx + 0.48, dy + 0.13, 0.10, BORDER_STRONG)
    tx(s, dx + 0.72, dy + 0.10, 4.0, 0.2, "neurolink.health  /  caregiver@…", size=7, color=MUTED)
    # sidebar
    rect(s, dx + 0.02, dy + 0.36, 1.62, dh - 0.38, fill=ACCENT, radius=False)
    tx(s, dx + 0.20, dy + 0.52, 1.3, 0.22, "NEUROLINK", size=8.5, color=LIGHT, bold=True, track=80)
    for i, item in enumerate(["Live Overview", "Alerts", "Safety", "Trends", "Management", "Settings"]):
        act = i == 0
        if act:
            rect(s, dx + 0.10, dy + 0.82 + i * 0.32, 1.44, 0.26, fill=ACCENT3, radius=True, adj=0.3)
        tx(s, dx + 0.22, dy + 0.85 + i * 0.32, 1.3, 0.2, item, size=7.5,
           color=LIGHT if act else LIGHT2, bold=act)
    dot(s, dx + 0.22, dy + 2.92, 0.08, OK)
    tx(s, dx + 0.36, dy + 2.84, 1.2, 0.24, "Band connected", size=7, color=LIGHT2)
    tx(s, dx + 0.20, dy + 3.62, 1.36, 0.4, "Margaret T. · 78\nCaregiver role", size=7, color=LIGHT2, line=1.25)
    # banner
    rect(s, dx + 1.72, dy + 0.44, dw - 1.80, 0.26, fill=OK_BG, radius=True, adj=0.3)
    tx(s, dx + 1.86, dy + 0.47, 5.4, 0.2, "● NLW-8842-A online · battery 68% · last sync 2 s ago", size=7.5, color=OK)
    # vitals cards
    vitals = [("Heart Rate", "72", "bpm", ACCENT), ("Blood Oxygen", "97", "%", ACCENT3),
              ("Skin Temp", "36.4", "°C", PURPLE), ("Stress Index", "0.28", "/1", ACCENT2)]
    vw = (dw - 1.88 - 0.24) / 4
    for i, (label, val, unit, col) in enumerate(vitals):
        vx = dx + 1.80 + i * (vw + 0.08)
        card(s, vx, dy + 0.80, vw, 0.78, fill=SURF2, line=BORDER)
        tx(s, vx + 0.12, dy + 0.88, vw - 0.3, 0.2, label, size=7, color=MUTED)
        rich(s, vx + 0.12, dy + 1.06, vw - 0.24, 0.30, [{"runs": [
            {"t": val, "size": 14, "bold": True, "color": col}, {"t": " " + unit, "size": 7, "color": MUTED}]}])
        spark(s, vx + 0.12, dy + 1.36, vw - 0.26, 0.16,
              [1, 2, 1.5, 2.4, 2, 1.7, 2.6, 2.1, 1.8, 2.2, 2.0, 1.9], color=col, lw=1.0)
    # charts
    card(s, dx + 1.80, dy + 1.70, 3.60, 1.36, fill=SURF2, line=BORDER)
    tx(s, dx + 1.92, dy + 1.80, 3.2, 0.2, "Heart-rate variability · 24 h", size=7.5, color=TEXT2, bold=True)
    spark(s, dx + 1.94, dy + 2.10, 3.32, 0.62,
          [52, 48, 55, 50, 46, 58, 61, 57, 49, 44, 47, 56, 62, 59, 51, 45, 48, 54], color=ACCENT3, lw=1.4)
    tx(s, dx + 1.92, dy + 2.80, 3.3, 0.2, "RMSSD (ms) · clinical overlay at 20 ms", size=6.5, color=MUTED)
    card(s, dx + 5.52, dy + 1.70, 2.02, 1.36, fill=SURF2, line=BORDER)
    tx(s, dx + 5.64, dy + 1.80, 1.8, 0.2, "Stress vs HR", size=7.5, color=TEXT2, bold=True)
    bars(s, dx + 5.66, dy + 2.14, 1.72, 0.56,
         [0.3, 0.42, 0.55, 0.38, 0.6, 0.72, 0.5, 0.44],
         colors=[PURPLE, PURPLE, WARN, PURPLE, WARN, DANGER, PURPLE, PURPLE], gap=0.06)
    tx(s, dx + 5.64, dy + 2.80, 1.9, 0.2, "threshold 0.60", size=6.5, color=MUTED)
    # alert list
    card(s, dx + 1.80, dy + 3.14, 5.74, 0.94, fill=SURF2, line=BORDER)
    tx(s, dx + 1.92, dy + 3.22, 3.0, 0.2, "RECENT ALERTS", size=7, color=MUTED, bold=True, track=60)
    alerts = [("Fall Detected", "critical", DANGER, DANGER_BG), ("High Stress", "high", WARN, WARN_BG),
              ("Fatigue marker", "low", INFO, INFO_BG)]
    ay = dy + 3.44
    for t, sev, col, bgc in alerts:
        chip(s, dx + 1.92, ay + 0.015, 0.86, 0.19, sev.upper(), fg=col, bgc=bgc, size=5.5, track=20)
        tx(s, dx + 2.86, ay, 2.6, 0.22, t, size=7.5, color=INK, bold=True)
        tx(s, dx + 5.30, ay, 1.6, 0.22, "AI explanation ready", size=6.5, color=MUTED, align="r")
        ay += 0.24

    # ── right: view list
    x = 8.62
    tx(s, x, 1.74, 3.92, 0.24, "THE SIX SHIPPED VIEWS", size=8, color=MUTED, bold=True, track=120)
    views = [
        ("Live Overview", "Vitals, IMU motion, live feed, device health, AI insight, quick actions."),
        ("Alerts", "Severity, evidence readings, explanation, recommendations, acknowledge / resolve."),
        ("Safety", "Incident timeline, GPS map, inactivity watch, one-click dispatch with delivery log."),
        ("Trends", "HRV, temperature & SpO₂, stress vs HR — 24 h / 48 h / 7 d / 30 d filters."),
        ("Management", "Contacts, devices & MQTT, thresholds, calibration, wearer profile, users."),
        ("Settings", "Broker configuration, simulator controls, theme, session and role handling."),
    ]
    y = 2.06
    for t, d in views:
        card(s, x, y, 3.92, 0.76, fill=SURF)
        tx(s, x + 0.18, y + 0.12, 3.6, 0.22, t, size=10.5, color=INK, bold=True)
        tx(s, x + 0.18, y + 0.36, 3.6, 0.34, d, size=8, color=TEXT2, line=1.22)
        y += 0.82
    footer(s, n, right="Product")


# ─────────────────────────────── 09 hardware ───────────────────────────
def hardware(b):
    s = slide(b.prs)
    n = b.num()
    header(s, "02 / Hardware", "Hardware: a five-sensor band on an ESP32",
           "Components referenced by the current firmware source. Measurements are prototype-grade "
           "and require bench validation.", num=n)

    rows = [["Component", "Measures", "Role in the system", "Interface"],
            ["MAX30102 — PPG", "Heart rate, SpO₂, HRV", "Cardiovascular context, oxygen and recovery signals", "I²C"],
            ["GSR electrodes", "Skin conductance (µS)", "Sympathetic arousal; tonic / phasic EDA split", "Analog"],
            ["MPU6050 — IMU", "Accel + gyro, 6 axes", "Activity state, tremor, steps, fall signature", "I²C"],
            ["MLX90614 — IR", "Skin temperature", "Fever trend and thermal drift, core-temperature estimate", "I²C (SMBus)"],
            ["NEO-6M — GNSS", "Latitude / longitude", "Location context attached to every escalated event", "UART"],
            ["ESP32-WROOM", "Edge compute + radio", "Feature extraction, threshold filtering, MQTT publish", "Wi-Fi / BLE"]]
    table(s, ML, 1.74, [2.20, 2.00, 5.60, 1.17], rows, row_h=0.44, head_h=0.34,
          size=9, head_size=8, align=["l", "l", "l", "l"], bold_col0=True)

    card(s, ML, 4.94, 6.62, 1.96, fill=SURF2)
    tx(s, ML + 0.26, 5.14, 6.1, 0.24, "FIRMWARE BEHAVIOUR (CURRENT SOURCE)", size=8, color=MUTED, bold=True, track=110)
    bullets(s, ML + 0.26, 5.46, 6.1, 1.3,
            [("Publishes JSON over MQTT ", "to neurolink/wear/<id>/telemetry every 1–5 seconds."),
             ("Runs the same equations as the server ", "— stress index, core-temperature estimate, fall physics."),
             ("Local threshold filter ", "so a broken link cannot silence a critical event."),
             ("Two OLED surfaces ", "for wearer-facing status and battery / link state.")],
            size=9, dot="·", gap=4, line=1.24)

    card(s, 7.68, 4.94, 4.86, 1.96, fill=SURF)
    tx(s, 7.94, 5.14, 4.3, 0.24, "PROTOTYPE → PRODUCT GATES", size=8, color=MUTED, bold=True, track=110)
    gates = [("Accuracy", "PPG and temperature validated against reference instruments"),
             ("Battery", "target multi-day wear between charges, measured per user"),
             ("Form factor", "wearable for 24 h including sleep, with a replaceable strap"),
             ("Manufacturing", "DFM review, component second-sourcing, EMC pre-scan")]
    gy = 5.46
    for t, d in gates:
        rich(s, 7.94, gy, 4.3, 0.34, [{"runs": [
            {"t": t + " — ", "size": 8.5, "bold": True, "color": ACCENT}, {"t": d, "size": 8.5, "color": TEXT2}], "line": 1.22}])
        gy += 0.34
    footer(s, n, right="Hardware")


# ─────────────────────────── 10 architecture ───────────────────────────
def architecture(b):
    s = slide(b.prs)
    n = b.num()
    header(s, "02 / Architecture", "One ingest path for simulated and real devices",
           "The end-to-end software path is implemented and running; device timing is the field-readiness gate.",
           num=n)

    def layer(x, y, w, h, title, sub, lines, col, fill=SURF, text_col=TEXT2):
        card(s, x, y, w, h, fill=fill, line=BORDER)
        rect(s, x, y, w, 0.05, fill=col, radius=True, adj=0.5)
        tx(s, x + 0.20, y + 0.18, w - 0.4, 0.24, title, size=10.5, color=INK, bold=True)
        if sub:
            tx(s, x + 0.20, y + 0.44, w - 0.4, 0.22, sub, size=7.5, color=col, bold=True, track=40)
        ly = y + 0.72
        for ln in lines:
            tx(s, x + 0.20, ly, w - 0.4, 0.24, ln, size=8.5, color=text_col)
            ly += 0.26

    def arrow(x, y, w=0.30, col=BORDER_STRONG):
        a = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y), Inches(w), Inches(0.17))
        a.shadow.inherit = False
        a.fill.solid(); a.fill.fore_color.rgb = rgb(col); a.line.fill.background()

    # row 1 — sources
    layer(ML, 1.72, 3.20, 1.52, "Sensor sources", "EDGE", ["ESP32 band · MQTT publish", "Real GNSS + IMU payloads",
                                                            "Simulator scenario cycle"], ACCENT)
    layer(ML + 3.46, 1.72, 3.20, 1.52, "Ingest & bridge", "SERVER", ["MQTT bridge (TLS :8883 / WS :8884)",
                                                                      "POST /api/telemetry/ingest",
                                                                      "Payload normaliser"], ACCENT3)
    layer(ML + 6.92, 1.72, 4.85, 1.52, "Intelligence core", "AI LAYER",
          ["Online per-wearer baselines + physiology equations", "Two-tier detection rules (clinical OR personal)",
           "Plain-language explanation + recommendations"], PURPLE)
    arrow(ML + 3.24, 2.40); arrow(ML + 6.70, 2.40)

    # row 2 — the pipeline ribbon
    card(s, ML, 3.42, CW, 0.62, fill=SURF2)
    tx(s, ML + 0.24, 3.52, 2.4, 0.24, "DATA SPINE", size=7.5, color=MUTED, bold=True, track=110)
    stages = ["Reading", "Baseline update", "Two-tier check", "Explain", "Persist", "Broadcast", "Caregiver view"]
    sx = ML + 1.62
    for i, st in enumerate(stages):
        w = 1.30
        tx(s, sx, 3.60, w, 0.24, st, size=8.5, color=INK, bold=True, align="c")
        if i < len(stages) - 1:
            tx(s, sx + w - 0.06, 3.60, 0.16, 0.24, "›", size=9.5, color=ACCENT3, align="c", bold=True)
        sx += w

    # row 3 — storage + api + clients
    layer(ML, 4.22, 3.20, 1.60, "Storage", "SQLITE", ["17 tables — vitals, alerts, contacts,",
                                                       "devices, thresholds, calibration refs,", "dispatch log, users"], OK)
    layer(ML + 3.46, 4.22, 3.20, 1.60, "Interfaces", "FASTAPI",
          ["REST /api/* (auth, telemetry, alerts,", "dispatch, devices, calibration, trends)",
           "WebSocket /ws · OpenAPI docs"], ACCENT3)
    layer(ML + 6.92, 4.22, 4.85, 1.60, "Clients", "WEB + MOBILE WEB",
          ["Responsive caregiver SPA (vanilla ES modules)", "No build step · dark / light themes",
           "Push-ready alert surfaces"], PURPLE)
    arrow(ML + 3.24, 4.90); arrow(ML + 6.70, 4.90)

    # row 4 — cross-cutting
    card(s, ML, 6.00, CW, 0.86, fill=SURF)
    tx(s, ML + 0.24, 6.14, 2.6, 0.24, "CROSS-CUTTING", size=7.5, color=MUTED, bold=True, track=110)
    cross = [("Roles & sessions", "PBKDF2 + HMAC, 30 days"), ("Observability", "staged connection tests"),
             ("Explainability", "every alert carries evidence"), ("Deployability", "$10 VPS to Pi"),
             ("Offline tolerance", "local narrative fallback")]
    cx = ML + 0.24
    for t, d in cross:
        rich(s, cx, 6.42, 2.28, 0.36, [{"runs": [
            {"t": t + " · ", "size": 8.5, "bold": True, "color": INK},
            {"t": d, "size": 8.5, "color": TEXT2}], "line": 1.2}])
        cx += 2.36
    footer(s, n, right="Architecture")
