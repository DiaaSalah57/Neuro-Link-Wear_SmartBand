"""
Deck QA — flags off-canvas shapes and real text collisions.

Text collisions are measured on the *ink* (wrapped glyph) extents rather than the
declared shape boxes, so wide-but-empty text boxes do not create false positives.

Run:  python3 presentation/qa_deck.py [deck.pptx]
"""
from __future__ import annotations

import sys
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn

sys.path.insert(0, str(Path(__file__).resolve().parent))
from preview import F, wrap, emu2px  # noqa: E402  (approx font metrics)

W_IN, H_IN = 13.333, 7.5
DPI = 100
EMU = 914400.0


def bbox(sh):
    return (sh.left / EMU, sh.top / EMU, (sh.left + sh.width) / EMU, (sh.top + sh.height) / EMU)


def ink_box(sh):
    """Approximate rendered text extents, in inches."""
    if sh.has_table:
        return bbox(sh)
    if not sh.has_text_frame:
        return None
    tf = sh.text_frame
    if not tf.text.strip():
        return None
    x0, y0, x1, y1 = bbox(sh)
    anchor = tf._txBody.find(qn("a:bodyPr")).get("anchor", "t")
    lines = []
    total_h = 0.0
    for p in tf.paragraphs:
        runs = [r for r in p.runs if r.text]
        if not runs:
            continue
        size = max((r.font.size.pt if r.font.size else 12) for r in runs)
        align = str(p.alignment) if p.alignment is not None else "LEFT"
        font = F(size, any(r.font.bold for r in runs))
        text = "".join(r.text for r in runs)
        # DejaVu (this checker) is ~12 % wider than Segoe UI (the deck font): wrap
        # at the wider effective width, then scale measured ink back down so the
        # extent stays comparable with the real rendering.
        SCALE = 0.88
        wrapped = wrap(text, font, ((x1 - x0) * DPI + 1) / SCALE)
        lh = size * (p.line_spacing if isinstance(p.line_spacing, float) else 1.0) * 1.20 / 72.0
        for ln in wrapped:
            w_in = font.getlength(ln) / DPI * SCALE
            if "CENTER" in align:
                lx = x0 + ((x1 - x0) - w_in) / 2
            elif "RIGHT" in align:
                lx = x1 - w_in
            else:
                lx = x0
            lines.append((lx, total_h, w_in, lh))
        total_h += lh * len(wrapped) + ((p.space_after.pt if p.space_after else 0) / 72.0)
    if anchor == "ctr":
        base = y0 + ((y1 - y0) - total_h) / 2
    elif anchor == "b":
        base = y1 - total_h
    else:
        base = y0
    xs0 = min(l[0] for l in lines)
    xs1 = max(l[0] + l[2] for l in lines)
    return (xs0, base, xs1, base + total_h)


def main(path):
    prs = Presentation(path)
    issues = 0
    for i, slide in enumerate(prs.slides, start=1):
        dark = None
        try:
            fill = slide.background.fill
            dark = str(fill.fore_color.rgb) == "0D0A24"
        except Exception:
            dark = False
        shapes = [sh for sh in slide.shapes if sh.shape_type != MSO_SHAPE_TYPE.LINE]
        for sh in shapes:
            x0, y0, x1, y1 = bbox(sh)
            if x0 < -0.02 or y0 < -0.02 or x1 > W_IN + 0.02 or y1 > H_IN + 0.02:
                print(f"[{i:02d}] off-canvas: {sh.name} ({x0:.2f},{y0:.2f})-({x1:.2f},{y1:.2f})")
                issues += 1
        boxes = []
        for sh in shapes:
            ib = ink_box(sh)
            if ib:
                label = sh.text_frame.text if sh.has_text_frame else sh.name
                boxes.append((ib, label.strip().replace("\n", " ")[:40]))
        for a in range(len(boxes)):
            for b in range(a + 1, len(boxes)):
                (ax0, ay0, ax1, ay1), ta = boxes[a]
                (bx0, by0, bx1, by1), tb = boxes[b]
                ox = min(ax1, bx1) - max(ax0, bx0)
                oy = min(ay1, by1) - max(ay0, by0)
                if ox > 0.05 and oy > 0.04:
                    print(f"[{i:02d}] text collision {ox*oy:.2f}in²  {ta!r} <> {tb!r}")
                    issues += 1
        if not dark:
            content = [sh for sh in shapes if ink_box(sh)]
            body = [sh for sh in content if bbox(sh)[1] < 6.9]
            if body:
                maxy = max(ink_box(sh)[3] for sh in body)
                if maxy < 6.05:
                    print(f"[{i:02d}] empty band: text ends at {maxy:.2f}in")
                    issues += 1
    print(f"\n{len(prs.slides)} slides checked · {issues} issue(s)")
    return 0 if issues == 0 else 1


if __name__ == "__main__":
    default = Path(__file__).resolve().parents[1] / "NeuroLink_Wear_Investor_Deck.pptx"
    sys.exit(main(sys.argv[1] if len(sys.argv) > 1 else str(default)))
